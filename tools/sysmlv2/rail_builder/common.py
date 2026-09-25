from __future__ import annotations

import hashlib
import json
import re
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook

TOOL_DIR = Path(__file__).resolve().parent
PROJECT_ROOT = TOOL_DIR.parents[2]
PRIMARY_EXCEL = PROJECT_ROOT / "Rail_MBSE_Function_Interface_v1_final.xlsx"
ARCH_JSON = PROJECT_ROOT / "work" / "architecture_xmi_ready_v1.json"
OUT_ROOT = PROJECT_ROOT / "work" / "sysmlv2" / "full_engineering_model"
GENERATED = OUT_ROOT / "01_generated"
MAPPING = OUT_ROOT / "02_mapping"
VALIDATION = OUT_ROOT / "03_validation"
REPORTS = OUT_ROOT / "04_reports"
SNAPSHOTS = OUT_ROOT / "05_snapshots"
AUDIT = OUT_ROOT / "00_source_audit"

SHEET_IDS = {
    "Function": "function_id",
    "Function_Allocation": "allocation_id",
    "Interface": "interface_instance_id",
    "Function_Interface": "mapping_id",
    "Connection": "connection_id",
    "PhysicalNet": "net_id",
    "Net_Member": "net_member_id",
    "Product_Hierarchy": "code",
    "Item": "item_code",
    "XMI_Export_QA": "qa_id",
}

SHEET_PURPOSE = {
    "Function": "Engineering functions and hierarchy level",
    "Function_Allocation": "Function-to-product allocations",
    "Interface": "Leaf raw port instances and item/type bindings",
    "Function_Interface": "Function-to-interface usage relations",
    "Connection": "Formal final leaf connections",
    "PhysicalNet": "N-ary physical network definitions",
    "Net_Member": "Physical network member ports",
    "Product_Hierarchy": "Complete L1-L4 product hierarchy",
    "Item": "Canonical engineering item dictionary",
    "XMI_Export_QA": "Source-side cross-check only",
}


def ensure_dirs() -> None:
    for path in (AUDIT, GENERATED, MAPPING, VALIDATION, REPORTS, SNAPSHOTS):
        path.mkdir(parents=True, exist_ok=True)


def read_tables(path: Path = PRIMARY_EXCEL) -> tuple[dict[str, list[dict]], list[dict]]:
    wb = load_workbook(path, read_only=True, data_only=True)
    tables: dict[str, list[dict]] = {}
    sheet_meta: list[dict] = []
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        id_col = SHEET_IDS.get(ws.title)
        header_idx = next(
            (i for i, row in enumerate(rows[:12]) if id_col and id_col in [str(v).strip() for v in row if v is not None]),
            None,
        )
        if header_idx is None:
            tables[ws.title] = []
            sheet_meta.append({"sheet": ws.title, "header_row": None, "rows": 0, "columns": [], "types": {}})
            continue
        headers = [str(v).strip() if v is not None else "" for v in rows[header_idx]]
        while headers and not headers[-1]:
            headers.pop()
        records = []
        for excel_row, row in enumerate(rows[header_idx + 1 :], start=header_idx + 2):
            values = list(row[: len(headers)])
            if not any(v is not None and str(v).strip() for v in values):
                continue
            rec = dict(zip(headers, values))
            rec["_source_file"] = path.name
            rec["_source_sheet"] = ws.title
            rec["_source_row"] = excel_row
            records.append(rec)
        tables[ws.title] = records
        types = {h: sorted({type(r[h]).__name__ for r in records if r.get(h) is not None}) for h in headers}
        ids = [str(r.get(id_col) or "").strip() for r in records]
        sheet_meta.append(
            {
                "sheet": ws.title,
                "header_row": header_idx + 1,
                "rows": len(records),
                "columns": headers,
                "types": types,
                "id_column": id_col,
                "blank_ids": sum(not value for value in ids),
                "duplicate_ids": sorted({value for value in ids if value and ids.count(value) > 1}),
                "purpose": SHEET_PURPOSE.get(ws.title, "Reference"),
                "participates": ws.title != "XMI_Export_QA",
            }
        )
    return tables, sheet_meta


def sid(prefix: str, value: object, hashed: bool = False) -> str:
    text = str(value or "")
    if hashed:
        return f"{prefix}_{hashlib.sha256(text.encode('utf-8')).hexdigest()[:16]}"
    clean = re.sub(r"[^A-Za-z0-9_]", "_", text).strip("_")
    clean = re.sub(r"_+", "_", clean)
    if not clean or clean[0].isdigit():
        clean = "N_" + clean
    return f"{prefix}_{clean}"


def q(value: object) -> str:
    text = "" if value is None else str(value)
    return '"' + text.replace("\\", "\\\\").replace('"', '\\"').replace("\r", " ").replace("\n", " ") + '"'


def bool_lit(value: object) -> str:
    return "true" if bool(value) else "false"


def build_model() -> dict:
    tables, sheet_meta = read_tables()
    products = tables["Product_Hierarchy"]
    items = tables["Item"]
    ports = tables["Interface"]
    connections = tables["Connection"]
    functions = tables["Function"]
    allocations = tables["Function_Allocation"]
    nets = tables["PhysicalNet"]
    net_members = tables["Net_Member"]
    function_interfaces = tables["Function_Interface"]

    product_by_code = {str(r["code"]): r for r in products}
    item_by_code = {r["item_code"]: r for r in items}
    port_by_id = {r["port_id"]: r for r in ports}
    children: dict[str, list[str]] = defaultdict(list)
    for p in products:
        if p.get("parent_code"):
            children[str(p["parent_code"])].append(str(p["code"]))
    for values in children.values():
        values.sort()

    ports_by_product: dict[str, list[dict]] = defaultdict(list)
    for port in ports:
        port["sysml_name"] = sid("rp", port["port_id"], True)
        ports_by_product[str(port["product_code"])].append(port)

    def chain(code: str) -> list[str]:
        result = [code]
        while product_by_code[result[-1]].get("parent_code"):
            result.append(str(product_by_code[result[-1]]["parent_code"]))
        return list(reversed(result))

    def product_path(code: str) -> str:
        return ".".join(sid("p", c) for c in chain(code))

    for p in products:
        code = str(p["code"])
        p["sysml_def"] = sid("P", code)
        p["sysml_usage"] = sid("p", code)
        p["sysml_path"] = product_path(code)

    item_is_physical = {
        code: any(str(p.get("port_category")) == "物理" for p in ports if p.get("item_code") == code)
        or str(item_by_code[code].get("datatype")) == "PhysicalConnector"
        for code in item_by_code
    }

    provenance = []
    arch = json.loads(ARCH_JSON.read_text(encoding="utf-8-sig"))
    arch_functions = {f["function_id"]: f for f in arch.get("functions", [])}
    for function in functions:
        supplemental = arch_functions.get(function["function_id"], {})
        for field in ("derived_from", "aggregated_from_functions"):
            function[field] = supplemental.get(field) or []
            if function[field]:
                provenance.append(
                    {
                        "Element_ID": function["function_id"], "Field": field, "Value": function[field],
                        "Source_File": ARCH_JSON.name, "Source_Sheet": "functions", "Source_Row": list(arch_functions).index(function["function_id"]) + 1,
                        "Source_Type": "ARCHITECTURE_JSON",
                    }
                )
    boundaries = arch.get("external_boundaries", [])
    boundary_by_id = {b["boundary_id"]: b for b in boundaries}
    boundary_ports: dict[tuple[str, str], dict] = {}
    for index, boundary in enumerate(boundaries, start=1):
        boundary["sysml_def"] = sid("B", boundary["boundary_id"])
        boundary["sysml_usage"] = sid("boundary", boundary["boundary_id"])
        for field in ("name", "domain"):
            provenance.append(
                {
                    "Element_ID": boundary["boundary_id"], "Field": field, "Value": boundary.get(field),
                    "Source_File": ARCH_JSON.name, "Source_Sheet": "external_boundaries", "Source_Row": index,
                    "Source_Type": "ARCHITECTURE_JSON",
                }
            )
    for c in connections:
        c["sysml_name"] = sid("c", c["connection_id"])
        c["source_path"] = product_path(str(c["source_product_code"])) + "." + port_by_id[c["source_port_id"]]["sysml_name"]
        target_code = str(c.get("target_product_code") or "")
        if target_code in product_by_code:
            c["target_path"] = product_path(target_code) + "." + port_by_id[c["target_port_id"]]["sysml_name"]
            c["target_effective_port_id"] = c["target_port_id"]
        else:
            key = (target_code, c["item_code"])
            if key not in boundary_ports:
                boundary_ports[key] = {
                    "boundary_id": target_code,
                    "item_code": c["item_code"],
                    "port_id": f"BPORT::{target_code}::{c['item_code']}",
                    "sysml_name": sid("bp", f"{target_code}|{c['item_code']}", True),
                }
            bp = boundary_ports[key]
            c["target_path"] = sid("boundary", target_code) + "." + bp["sysml_name"]
            c["target_effective_port_id"] = bp["port_id"]
            provenance.append(
                {
                    "Element_ID": c["connection_id"], "Field": "Target_Raw_Port_ID", "Value": bp["port_id"],
                    "Source_File": PRIMARY_EXCEL.name, "Source_Sheet": "Connection", "Source_Row": c["_source_row"],
                    "Source_Type": "DERIVED",
                }
            )
        source_port = port_by_id[c["source_port_id"]]
        target_port = port_by_id.get(c.get("target_port_id"))
        c["physical_or_signal"] = "PHYSICAL" if item_is_physical[c["item_code"]] else "SIGNAL"
        candidates = {p.get("physical_net_id") for p in (source_port, target_port) if p and p.get("physical_net_id")}
        c["physical_net_id"] = sorted(candidates)[0] if len(candidates) == 1 else None
        c["original_path"] = f"{c['source_product_code']}::{c['source_port_id']} -> {c['target_product_code']}::{c['target_effective_port_id']}"
        for field, value, source_type in (
            ("Physical_or_Signal", c["physical_or_signal"], "DERIVED"),
            ("Physical_Net_ID", c["physical_net_id"], "DERIVED"),
            ("Original_Path", c["original_path"], "DERIVED"),
        ):
            provenance.append(
                {
                    "Element_ID": c["connection_id"], "Field": field, "Value": value,
                    "Source_File": PRIMARY_EXCEL.name, "Source_Sheet": "Connection", "Source_Row": c["_source_row"],
                    "Source_Type": source_type,
                }
            )

    return {
        "tables": tables, "sheet_meta": sheet_meta, "products": products, "product_by_code": product_by_code,
        "children": dict(children), "items": items, "item_by_code": item_by_code, "item_is_physical": item_is_physical,
        "ports": ports, "port_by_id": port_by_id, "ports_by_product": dict(ports_by_product),
        "connections": connections, "functions": functions, "allocations": allocations,
        "function_interfaces": function_interfaces, "nets": nets, "net_members": net_members,
        "boundaries": boundaries, "boundary_by_id": boundary_by_id, "boundary_ports": list(boundary_ports.values()),
        "provenance": provenance,
    }


def write_text(path: Path, text: str) -> None:
    path.write_text(text.rstrip() + "\n", encoding="utf-8", newline="\n")


def dump_json(path: Path, data: object) -> None:
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2, default=str) + "\n", encoding="utf-8")
