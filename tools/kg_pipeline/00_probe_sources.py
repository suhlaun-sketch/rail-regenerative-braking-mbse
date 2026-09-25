"""Inventory workbook sheets and locate D/E/X hierarchy evidence without modifying inputs."""

from __future__ import annotations

from pathlib import Path

from common import PROJECT_DIR, WORK_DIR, normalize_code, normalize_text, write_json


def main() -> int:
    from openpyxl import load_workbook

    payload = []
    for path in sorted(PROJECT_DIR.glob("*.xlsx")):
        if path.name.startswith("~$"):
            continue
        workbook_entry = {"path": str(path), "sheets": [], "dex_rows": []}
        try:
            wb = load_workbook(path, read_only=True, data_only=True)
        except Exception as exc:
            workbook_entry["error"] = str(exc)
            payload.append(workbook_entry)
            continue
        for ws in wb.worksheets:
            headers = [normalize_text(cell.value) for cell in ws[1]] if (ws.max_row or 0) else []
            workbook_entry["sheets"].append(
                {"name": ws.title, "max_row": ws.max_row or 0, "max_column": ws.max_column or 0, "headers": headers}
            )
            if "code" not in headers:
                continue
            code_index = headers.index("code")
            for row_number, values in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
                code = normalize_code(values[code_index] if code_index < len(values) else None)
                if code.startswith(("D", "E", "X")):
                    workbook_entry["dex_rows"].append(
                        {
                            "sheet": ws.title,
                            "row": row_number,
                            "values": [normalize_text(value) for value in values],
                        }
                    )
        wb.close()
        payload.append(workbook_entry)
    output = WORK_DIR / "source_probe.json"
    write_json(output, payload)
    print(f"Scanned workbooks: {len(payload)}")
    for entry in payload:
        print(f"- {Path(entry['path']).name}: D/E/X rows={len(entry.get('dex_rows', []))}")
    print(f"Probe: {output.resolve()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
