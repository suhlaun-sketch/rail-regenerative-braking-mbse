"""Extract approved Interface Layer into deterministic Neo4j load data."""
from __future__ import annotations
import argparse
from collections import defaultdict
from openpyxl import load_workbook
from common import *

def norm(v): return normalize_text(v) if v is not None else ""
def records(ws):
    hs=[norm(c.value) for c in ws[1]]; out=[]
    for rn, vals in enumerate(ws.iter_rows(min_row=2, values_only=True),2):
        if any(norm(v) for v in vals): out.append({h:(vals[i] if i<len(vals) else None) for i,h in enumerate(hs)}|{"_row":rn})
    return out
def val(r,key,default=""): return norm(r.get(key,default))

def extract_products(full, selected, cfg, pid):
    standard={}; by_name=defaultdict(list)
    for r in full:
        code=normalize_code(r.get("code"));
        if not code: continue
        if code in standard: raise ValueError(f"duplicate standard Product code: {code}")
        x=dict(r); x["code"]=code; x["name"]=val(r,"name"); x["parent_name"]=val(r,"parent_name"); standard[code]=x; by_name[x["name"]].append(x)
    custom={}
    for c in cfg.get("custom_products",[]):
        code=normalize_code(c["code"])
        if code in custom: raise ValueError(f"duplicate custom_products code: {code}")
        z=dict(c); z.update(code=code,name=norm(c.get("name")),parent_name=norm(c.get("parent_name")),parent_code=normalize_code(c.get("parent_code")) if c.get("parent_code") else "",level=int(c["level"]),source_type=norm(c.get("source_type","project_extension")),source_ref=norm(c.get("source_ref",f"CONFIG:custom_products:{code}")))
        custom[code]=z
    nodes={**standard}
    for code,c in custom.items():
        if code not in nodes: nodes[code]={"code":code,"name":c["name"],"parent_name":c["parent_name"],"角色标签":c.get("node_role",""),"_row":0}
    sel={}
    for r in selected:
        code=normalize_code(r.get("code"));
        if not code: continue
        if code in sel: raise ValueError(f"duplicate selected Product code: {code}")
        if code not in nodes: raise ValueError(f"selected Product missing from standard source and custom_products: {code}")
        sel[code]=r
    closure=set(); parents={}
    for leaf in sel:
        cur=leaf; seen=set()
        while cur:
            if cur in seen: raise ValueError(f"Product hierarchy cycle at {cur}")
            seen.add(cur); closure.add(cur); n=nodes[cur]; c=custom.get(cur); level=int(c["level"]) if c else product_level(cur)
            if level==1:
                if c and c.get("parent_code"): raise ValueError(f"level1 custom Product has parent: {cur}")
                # Some standard source rows contain a stray legacy value in parent_name;
                # level=1 is authoritative and such a value must not create a fake parent.
                break
            if c and c.get("parent_code"):
                p=c["parent_code"]
                if p not in nodes: raise ValueError(f"custom parent missing: {cur}->{p}")
                if c.get("parent_name") and nodes[p]["name"]!=c["parent_name"]: raise ValueError(f"custom parent name mismatch: {cur}")
            else:
                expected=expected_parent_code(cur); p=expected if expected in nodes and nodes[expected]["name"]==val(n,"parent_name") else ""
                if not p:
                    cand=[q for q in by_name.get(val(n,"parent_name"),[]) if product_level(q["code"])==level-1]
                    if len(cand)!=1: raise ValueError(f"cannot resolve parent: {cur} parent_name={val(n,'parent_name')!r}")
                    p=cand[0]["code"]
            parents[cur]=p; cur=p
    products=[]
    for code in sorted(closure,key=lambda x:(int(custom.get(x,{}).get("level",product_level(x))),x)):
        n=nodes[code]; c=custom.get(code,{}); sr=sel.get(code); pc=parents.get(code,""); pn=nodes[pc]["name"] if pc else ""
        products.append({"code":code,"name":c.get("name",n["name"]),"level":int(c.get("level",product_level(code))),"parent_code":pc,"parent_name":pn,"selected":True,"leaf":code in sel,"node_role":norm((sr or {}).get("节点角色")) or norm(c.get("node_role",n.get("角色标签",""))),"namespace":namespace_for_code(code),"project_id":pid,"source_type":c.get("source_type","standard_299"),"source_ref":c.get("source_ref",f"全部!{n.get('_row',0)}"),"source_sheet":"入选" if sr else ("CONFIG:custom_products" if c else "全部"),"source_row":int((sr or {}).get("_row",n.get("_row",0))),"source_basis":norm(c.get("source_basis",""))})
    contains=[{"parent_code":p["parent_code"],"child_code":p["code"],"project_id":pid,"relation_id":f"CONTAINS::{p['parent_code']}::{p['code']}"} for p in products if p["parent_code"]]
    return products,contains

def extract_items(rows, used, pid):
    out=[]
    for r in rows:
        code=normalize_code(r.get("item_code"));
        if not code: continue
        vals=[r.get(k) for k in list(r) if not k.startswith("_")]+[None]*13
        out.append({"item_code":code,"item_name_cn":norm(vals[1]),"item_name_en":norm(vals[2]),"item_family":norm(vals[3]),"domain":norm(vals[4]),"datatype":norm(vals[5]),"unit_medium":norm(vals[6]),"semantic_definition":norm(vals[7]),"aliases":norm(vals[8]),"priority":norm(vals[9]),"validity_rule":norm(vals[10]),"fault_strategy":norm(vals[11]),"usage_note":norm(vals[12]),"port_type_id":stable_port_type_id(code),"interface_type_id":stable_interface_type_id(code),"used_in_current_architecture":code in used,"project_id":pid})
    if len(out)!=144: raise ValueError(f"Item字典应为144条，实际{len(out)}")
    return out

def main():
    argparse.ArgumentParser().parse_args(); cfg=load_config(); ensure_inputs()
    wb=load_workbook(INPUT_WORKBOOK,data_only=True); iw=load_workbook(ITEM_WORKBOOK,data_only=True)
    full,selected,raw=records(wb["全部"]),records(wb["入选"]),records(wb["Sheet3"]); itemrows=records(iw["Item字典"])
    products,contains=extract_products(full,selected,cfg,cfg["project_id"]); ptmap=records(wb["PortType映射"]); used={normalize_code(r.get("item_code")) for r in ptmap if r.get("item_code")}; items=extract_items(itemrows,used,cfg["project_id"])
    item_by={x["item_name_cn"]:x for x in items}; leaf={x["code"] for x in products if x["leaf"]}; profile=cfg["profile"]; exposes=[]
    inactive_dc_products={"4121","4212","5121","5122"}
    for r in raw:
        code=normalize_code(r.get("code")); item=item_by.get(val(r,"交换内容"));
        if code not in leaf: raise ValueError(f"Raw Port product is not selected leaf: {code}")
        if not item: raise ValueError(f"Raw Port item not canonical: {val(r,'交换内容')}")
        direction=val(r,"方向"); profile_active=not (code in inactive_dc_products and item["item_code"]=="ITM-PHY-002"); exposes.append({"port_id":stable_port_id(code,val(r,"端口名称"),item["item_code"],direction),"product_code":code,"item_code":item["item_code"],"port_name":val(r,"端口名称"),"direction":direction,"port_category":val(r,"端口类别"),"unit_medium":val(r,"单位/介质"),"port_type_id":stable_port_type_id(item["item_code"]),"interface_type_id":stable_interface_type_id(item["item_code"]),"profile_id":profile["profile_id"],"profile_active":profile_active,"source_code":code,"source_sheet":"Sheet3","source_row":int(r["_row"]),"project_id":cfg["project_id"]})
    active=sum(1 for x in exposes if x["profile_active"]); inactive=len(exposes)-active
    data={"meta":{"project_id":cfg["project_id"],"profile":profile,"counts":{"selected_leaves":len(selected),"products":len(products),"items":len(items),"contains":len(contains),"exposes":len(exposes),"active_exposes":active,"inactive_exposes":inactive},"warnings":["D/X are project_extension; E100 is external_system level-2 leaf."]},"products":products,"contains":contains,"items":items,"exposes":exposes}; write_json(GRAPH_DATA_PATH,data); print(data["meta"]["counts"])
if __name__=="__main__": main()
