#!/usr/bin/env python3
"""Build deterministic L3/L2/L1 interface aggregation artifacts.

Inputs are the frozen 147 Product / 184 Final Connection / 517 Raw Port /
144 Item Dictionary data already integrated in ``work/architecture_xmi_ready_v1.json``
and the authoritative Final Connection list.  This script does not modify any
input artifact.
"""

from __future__ import annotations

import hashlib
import json
import os
import sys
from collections import Counter, defaultdict
from copy import copy
from datetime import datetime
from pathlib import Path
from typing import Any, Iterable


ROOT = Path(__file__).resolve().parents[2]

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.formatting.rule import FormulaRule
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    # The project keeps its reproducible Python dependencies in this venv.
    local_site = ROOT / "tools" / "interface_pipeline" / ".venv" / "Lib" / "site-packages"
    if local_site.exists():
        sys.path.insert(0, str(local_site))
        from openpyxl import Workbook, load_workbook
        from openpyxl.formatting.rule import FormulaRule
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter
    else:
        raise


ARCHITECTURE_PATH = ROOT / "work" / "architecture_xmi_ready_v1.json"
FINAL_CONNECTIONS_PATH = ROOT / "tools" / "topology_pipeline" / "work" / "final_connections.json"
UPWARD_PORTS_PATH = ROOT / "tools" / "xmi_pipeline" / "work" / "upward_port_inference_v1.json"
HIERARCHICAL_PROJECTION_PATH = (
    ROOT / "tools" / "xmi_pipeline" / "work" / "hierarchical_connection_projection_v1.json"
)
OUTPUT_XLSX = ROOT / "work" / "Rail_MBSE_分层接口聚合_v1.xlsx"
OUTPUT_JSON = ROOT / "work" / "Rail_MBSE_分层接口聚合_v1.json"

EXPECTED_COUNTS = {"products": 147, "items": 144, "interfaces": 517, "connections": 184}
LEVELS = (("L3", 3, "01_三级聚合"), ("L2", 2, "02_二级聚合"), ("L1", 1, "03_一级聚合"))
FIXED_EXCEL_TIME = datetime(2000, 1, 1, 0, 0, 0)

COLUMNS = [
    "Aggregate_Interface_ID",
    "Aggregate_Level",
    "Boundary_Type",
    "Source_Aggregate_Code",
    "Source_Aggregate_Name",
    "Source_Aggregate_Port_Name",
    "Source_Aggregate_Direction",
    "Target_Aggregate_Code",
    "Target_Aggregate_Name",
    "Target_Aggregate_Port_Name",
    "Target_Aggregate_Direction",
    "Item_Code",
    "Item_Name",
    "Item_Name_EN",
    "Item_Family",
    "Physical_or_Signal",
    "Port_Category",
    "Domain",
    "Datatype",
    "Unit_or_Medium",
    "Semantic_Definition",
    "Aliases",
    "Usage_Note",
    "Port_Type_ID",
    "Interface_Type_ID",
    "Signal_Role",
    "Energy_Role",
    "Priority",
    "Validity_Rule",
    "Fault_Tolerance_Strategy",
    "Condition_Mode",
    "Safety_Related",
    "Flow_Key",
    "Leaf_Connection_Count",
    "Original_Connection_IDs",
    "Source_Leaf_Codes",
    "Source_Leaf_Names",
    "Source_Raw_Port_IDs",
    "Source_Raw_Port_Names",
    "Source_Raw_Directions",
    "Target_Leaf_Codes",
    "Target_Leaf_Names",
    "Target_Raw_Port_IDs",
    "Target_Raw_Port_Names",
    "Target_Raw_Directions",
    "Original_Connection_Types",
    "Resolution_Basis",
    "Original_Connection_Status",
    "Profile",
    "Active_Status",
    "Original_Path_Summary",
    "Trace_Status",
    "Attribute_Propagation_Status",
    "Notes",
]


def load_json(path: Path) -> dict[str, Any]:
    with path.open("r", encoding="utf-8") as handle:
        return json.load(handle)


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def canonical_digest(value: Any) -> str:
    payload = json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    return hashlib.sha256(payload.encode("utf-8")).hexdigest()


def stable_values(values: Iterable[Any]) -> list[str]:
    return sorted({str(value) for value in values if value is not None and str(value) != ""})


def stable_join(values: Iterable[Any]) -> str:
    return " | ".join(stable_values(values))


def aligned_join(values: Iterable[Any]) -> str:
    """Join values without sorting/deduplicating, preserving blank positions."""
    return " | ".join("" if value is None else str(value) for value in values)


def active_status(values: Iterable[Any]) -> str:
    normalized = {bool(value) for value in values if value is not None}
    if normalized == {True}:
        return "ACTIVE"
    if normalized == {False}:
        return "INACTIVE"
    if normalized == {True, False}:
        return "MIXED"
    return ""


def flow_key(item: dict[str, Any], physical_or_signal: str) -> str:
    base = [
        item.get("item_code", ""),
        item.get("domain", ""),
        physical_or_signal.lower(),
    ]
    if physical_or_signal == "Signal":
        base.extend([item.get("datatype", ""), item.get("unit_medium", "")])
    elif item.get("unit_medium"):
        base.append(item["unit_medium"])
    return "|".join(str(value) for value in base)


def main() -> int:
    architecture = load_json(ARCHITECTURE_PATH)
    final_payload = load_json(FINAL_CONNECTIONS_PATH)
    upward = load_json(UPWARD_PORTS_PATH)
    hierarchical_projection = load_json(HIERARCHICAL_PROJECTION_PATH)

    actual_counts = {
        "products": len(architecture.get("products", [])),
        "items": len(architecture.get("items", [])),
        "interfaces": len(architecture.get("interfaces", [])),
        "connections": len(final_payload.get("connections", [])),
    }
    if actual_counts != EXPECTED_COUNTS:
        raise RuntimeError(f"Frozen input count mismatch: expected {EXPECTED_COUNTS}, got {actual_counts}")
    if upward.get("raw_port_count") != 517 or upward.get("promoted_port_count") != 594:
        raise RuntimeError("Upward inference freeze mismatch for Raw/Promoted Ports")
    if hierarchical_projection.get("hierarchical_segment_count") != 724:
        raise RuntimeError("Upward inference freeze mismatch for hierarchical connector segments")
    if len(hierarchical_projection.get("segments", [])) != 724:
        raise RuntimeError("Hierarchical projection segment list is incomplete")

    products = architecture["products"]
    items = architecture["items"]
    interfaces = architecture["interfaces"]
    final_connections = final_payload["connections"]
    boundaries = architecture.get("external_boundaries", [])

    product_by_code = {row["id"]: row for row in products}
    item_by_code = {row["item_code"]: row for row in items}
    port_by_id = {row["port_id"]: row for row in interfaces}
    boundary_by_code = {row["boundary_id"]: row for row in boundaries}
    connection_ids = [row["connection_id"] for row in final_connections]
    if len(product_by_code) != 147 or len(item_by_code) != 144 or len(port_by_id) != 517:
        raise RuntimeError("Frozen Product/Item/Raw Port identifiers are not unique")
    if len(set(connection_ids)) != 184:
        raise RuntimeError("Frozen Final Connection identifiers are not unique")

    architecture_connection_ids = {row["connection_id"] for row in architecture["connections"]}
    if architecture_connection_ids != set(connection_ids):
        raise RuntimeError("Architecture and authoritative Final Connection sets differ")

    def ancestor(code: str | None, target_level: int) -> dict[str, Any] | None:
        if not code or code not in product_by_code:
            return None
        current = product_by_code[code]
        seen: set[str] = set()
        while int(current["level"]) > target_level:
            if current["id"] in seen:
                raise RuntimeError(f"Product hierarchy cycle at {current['id']}")
            seen.add(current["id"])
            parent_code = current.get("parent_id")
            if not parent_code or parent_code not in product_by_code:
                return None
            current = product_by_code[parent_code]
        # A shallow real leaf such as E100 remains itself; no fake L3/L4 is created.
        return current

    missing_item_ids: set[str] = set()
    missing_source_ancestor_ids: set[str] = set()
    missing_target_ancestor_ids: set[str] = set()
    unresolved_trace_ids: set[str] = set()
    level_records: dict[str, list[dict[str, Any]]] = {}

    for level_code, level_number, _sheet_name in LEVELS:
        groups: dict[tuple[str, str, str, str], list[dict[str, Any]]] = defaultdict(list)

        for connection in final_connections:
            connection_id = connection["connection_id"]
            item = item_by_code.get(connection.get("item_code"))
            source_product = product_by_code.get(connection.get("source_product"))
            source_port = port_by_id.get(connection.get("source_port_id"))
            source_aggregate = ancestor(connection.get("source_product"), level_number)

            target_is_boundary = not connection.get("target_product") and bool(connection.get("target_boundary"))
            target_product = product_by_code.get(connection.get("target_product"))
            target_port = port_by_id.get(connection.get("target_port_id")) if connection.get("target_port_id") else None
            if target_is_boundary:
                target_boundary = boundary_by_code.get(connection.get("target_boundary"))
                target_aggregate_code = connection.get("target_boundary", "")
                target_aggregate_name = (target_boundary or {}).get("name", target_aggregate_code)
            else:
                target_boundary = None
                target_aggregate = ancestor(connection.get("target_product"), level_number)
                target_aggregate_code = (target_aggregate or {}).get("id", "")
                target_aggregate_name = (target_aggregate or {}).get("name", "")

            if item is None:
                missing_item_ids.add(connection_id)
            if source_aggregate is None:
                missing_source_ancestor_ids.add(connection_id)
            if not target_is_boundary and not target_aggregate_code:
                missing_target_ancestor_ids.add(connection_id)
            if target_is_boundary and target_boundary is None:
                missing_target_ancestor_ids.add(connection_id)

            source_aggregate_code = (source_aggregate or {}).get("id", "")
            source_aggregate_name = (source_aggregate or {}).get("name", "")
            physical_or_signal = "Signal" if connection.get("mode") == "SIGNAL" else "Physical"
            key = (
                source_aggregate_code,
                target_aggregate_code,
                connection.get("item_code", ""),
                physical_or_signal,
            )

            target_leaf_code = connection.get("target_product") or connection.get("target_boundary") or ""
            target_leaf_name = (
                (target_product or {}).get("name", "")
                if connection.get("target_product")
                else (target_boundary or {}).get("name", target_leaf_code)
            )
            resolved = bool(
                item
                and source_product
                and source_port
                and source_aggregate
                and (
                    (target_is_boundary and target_boundary)
                    or (target_product and target_port and target_aggregate_code)
                )
            )
            if not resolved:
                unresolved_trace_ids.add(connection_id)

            path_summary = (
                f"{connection.get('source_product', '')}({(source_product or {}).get('name', '')})/"
                f"{(source_port or {}).get('name', '')} -> "
                f"{target_leaf_code}({target_leaf_name})/"
                f"{(target_port or {}).get('name', 'EXTERNAL_BOUNDARY' if target_is_boundary else '')} "
                f"[{connection.get('item_code', '')}:{(item or {}).get('name_cn', '')}]"
            )
            groups[key].append(
                {
                    "connection": connection,
                    "item": item or {},
                    "source_product": source_product or {},
                    "source_port": source_port or {},
                    "source_aggregate": source_aggregate or {},
                    "target_product": target_product or {},
                    "target_port": target_port or {},
                    "target_aggregate_code": target_aggregate_code,
                    "target_aggregate_name": target_aggregate_name,
                    "target_leaf_code": target_leaf_code,
                    "target_leaf_name": target_leaf_name,
                    "physical_or_signal": physical_or_signal,
                    "resolved": resolved,
                    "path_summary": path_summary,
                }
            )

        projected: list[dict[str, Any]] = []
        for key, members in groups.items():
            source_code, target_code, item_code, physical_or_signal = key
            members.sort(key=lambda row: row["connection"]["connection_id"])
            first = members[0]
            item = first["item"]
            boundary_type = "INTERNAL" if source_code == target_code else "CROSS_BOUNDARY"
            if boundary_type == "INTERNAL":
                source_aggregate_port_name = "INTERNAL"
                target_aggregate_port_name = "INTERNAL"
                source_aggregate_direction = ""
                target_aggregate_direction = ""
                notes = "当前聚合层内部连接；未生成上层外部端口。"
            elif physical_or_signal == "Signal":
                source_aggregate_port_name = f"{item.get('name_cn', '')}输出"
                target_aggregate_port_name = f"{item.get('name_cn', '')}输入"
                source_aggregate_direction = "out"
                target_aggregate_direction = "in"
                notes = ""
            else:
                source_aggregate_port_name = f"{item.get('name_cn', '')}接口"
                target_aggregate_port_name = f"{item.get('name_cn', '')}接口"
                source_aggregate_direction = "inout"
                target_aggregate_direction = "inout"
                notes = ""

            source_categories = [row["source_port"].get("category") for row in members]
            target_categories = [row["target_port"].get("category") for row in members]
            trace_ok = all(row["resolved"] for row in members)

            port_type_id = stable_join(
                [item.get("port_type_id")]
                + [row["source_port"].get("port_type_id") for row in members]
                + [row["target_port"].get("port_type_id") for row in members]
            )
            interface_type_id = stable_join(
                [item.get("interface_type_id")]
                + [row["connection"].get("interface_type_id") for row in members]
                + [row["source_port"].get("interface_type_id") for row in members]
                + [row["target_port"].get("interface_type_id") for row in members]
            )
            resolution_basis = aligned_join(
                row["connection"].get("resolution_basis", "") for row in members
            )
            propagated = {
                "Item_Code": item_code,
                "Item_Name": item.get("name_cn", ""),
                "Item_Name_EN": item.get("name_en", ""),
                "Item_Family": item.get("family", ""),
                "Physical_or_Signal": physical_or_signal,
                "Port_Category": stable_join(source_categories + target_categories),
                "Domain": item.get("domain", ""),
                "Datatype": item.get("datatype", ""),
                "Unit_or_Medium": item.get("unit_medium", ""),
                "Semantic_Definition": item.get("semantic_definition", ""),
                "Aliases": item.get("aliases", ""),
                "Usage_Note": item.get("usage_note", ""),
                "Port_Type_ID": port_type_id,
                "Interface_Type_ID": interface_type_id,
                # These attributes are absent from the frozen Item/Connection schema.
                "Signal_Role": item.get("signal_role", ""),
                "Energy_Role": item.get("energy_role", ""),
                "Priority": item.get("priority", ""),
                "Validity_Rule": item.get("validity_rule", ""),
                "Fault_Tolerance_Strategy": item.get("fault_strategy", ""),
                "Condition_Mode": item.get("condition_mode", ""),
                "Safety_Related": item.get("safety_related", ""),
            }
            attribute_ok = bool(
                item
                and propagated["Item_Code"]
                and propagated["Item_Name"]
                and propagated["Domain"]
                and propagated["Physical_or_Signal"]
                and propagated["Unit_or_Medium"] == item.get("unit_medium", "")
                and propagated["Datatype"] == item.get("datatype", "")
                and propagated["Semantic_Definition"] == item.get("semantic_definition", "")
                and propagated["Aliases"] == item.get("aliases", "")
                and propagated["Usage_Note"] == item.get("usage_note", "")
                and propagated["Port_Type_ID"] == port_type_id
                and propagated["Interface_Type_ID"] == interface_type_id
            )

            record = {
                "Aggregate_Interface_ID": "",  # assigned after deterministic sorting
                "Aggregate_Level": level_code,
                "Boundary_Type": boundary_type,
                "Source_Aggregate_Code": source_code,
                "Source_Aggregate_Name": first["source_aggregate"].get("name", ""),
                "Source_Aggregate_Port_Name": source_aggregate_port_name,
                "Source_Aggregate_Direction": source_aggregate_direction,
                "Target_Aggregate_Code": target_code,
                "Target_Aggregate_Name": first["target_aggregate_name"],
                "Target_Aggregate_Port_Name": target_aggregate_port_name,
                "Target_Aggregate_Direction": target_aggregate_direction,
                **propagated,
                "Flow_Key": flow_key(item, physical_or_signal),
                "Leaf_Connection_Count": len(members),
                "Original_Connection_IDs": aligned_join(
                    row["connection"]["connection_id"] for row in members
                ),
                "Source_Leaf_Codes": aligned_join(
                    row["connection"].get("source_product") for row in members
                ),
                "Source_Leaf_Names": aligned_join(row["source_product"].get("name") for row in members),
                "Source_Raw_Port_IDs": aligned_join(
                    row["connection"].get("source_port_id") for row in members
                ),
                "Source_Raw_Port_Names": aligned_join(row["source_port"].get("name") for row in members),
                "Source_Raw_Directions": aligned_join(
                    row["source_port"].get("direction") for row in members
                ),
                "Target_Leaf_Codes": aligned_join(row["target_leaf_code"] for row in members),
                "Target_Leaf_Names": aligned_join(row["target_leaf_name"] for row in members),
                "Target_Raw_Port_IDs": aligned_join(
                    row["connection"].get("target_port_id") for row in members
                ),
                "Target_Raw_Port_Names": aligned_join(row["target_port"].get("name") for row in members),
                "Target_Raw_Directions": aligned_join(
                    row["target_port"].get("direction") for row in members
                ),
                "Original_Connection_Types": aligned_join(
                    row["connection"].get("mode") for row in members
                ),
                "Resolution_Basis": resolution_basis,
                # The frozen Connection schema has no row-level status attribute.
                "Original_Connection_Status": "",
                "Profile": aligned_join(row["connection"].get("profile_id") for row in members),
                "Active_Status": active_status(
                    value
                    for row in members
                    for value in (
                        row["source_port"].get("profile_active"),
                        row["target_port"].get("profile_active")
                        if row["target_port"]
                        else None,
                    )
                ),
                "Original_Path_Summary": aligned_join(row["path_summary"] for row in members),
                "Trace_Status": "PASS" if trace_ok else "FAIL",
                "Attribute_Propagation_Status": "PASS" if attribute_ok else "FAIL",
                "Notes": notes,
            }
            projected.append(record)

        projected.sort(
            key=lambda row: (
                0 if row["Boundary_Type"] == "CROSS_BOUNDARY" else 1,
                row["Source_Aggregate_Code"],
                row["Target_Aggregate_Code"],
                row["Flow_Key"],
                row["Original_Connection_IDs"],
            )
        )
        for index, record in enumerate(projected, 1):
            record["Aggregate_Interface_ID"] = f"AGG-{level_code}-{index:04d}"
        level_records[level_code] = projected

    # Verify every Final Connection is covered at all three aggregation levels.
    expected_connection_ids = set(connection_ids)
    coverage_by_level: dict[str, set[str]] = {}
    for level_code, _level_number, _sheet_name in LEVELS:
        covered: set[str] = set()
        for record in level_records[level_code]:
            covered.update(record["Original_Connection_IDs"].split(" | "))
        coverage_by_level[level_code] = covered
    covered_all_levels = set.intersection(*(coverage_by_level[level] for level, _, _ in LEVELS))

    # Verify that every trace column is aligned to the atomic connection-ID order.
    connection_by_id = {row["connection_id"]: row for row in final_connections}
    pair_alignment_errors: list[dict[str, Any]] = []
    paired_columns = [
        "Source_Leaf_Codes",
        "Source_Leaf_Names",
        "Source_Raw_Port_IDs",
        "Source_Raw_Port_Names",
        "Source_Raw_Directions",
        "Target_Leaf_Codes",
        "Target_Leaf_Names",
        "Target_Raw_Port_IDs",
        "Target_Raw_Port_Names",
        "Target_Raw_Directions",
        "Original_Connection_Types",
        "Resolution_Basis",
        "Profile",
        "Original_Path_Summary",
    ]
    for level_code, records in level_records.items():
        for record in records:
            ids = record["Original_Connection_IDs"].split(" | ")
            actual = {column: record[column].split(" | ") for column in paired_columns}
            bad_fields = [column for column in paired_columns if len(actual[column]) != len(ids)]
            if ids != sorted(ids):
                bad_fields.append("Original_Connection_IDs_ORDER")
            if not bad_fields:
                for index, connection_id in enumerate(ids):
                    connection = connection_by_id.get(connection_id)
                    if connection is None:
                        bad_fields.append("Original_Connection_IDs_RESOLUTION")
                        break
                    source_product = product_by_code[connection["source_product"]]
                    source_port = port_by_id[connection["source_port_id"]]
                    target_code = (
                        connection.get("target_product") or connection.get("target_boundary") or ""
                    )
                    target_product = product_by_code.get(connection.get("target_product") or "", {})
                    target_boundary = boundary_by_code.get(connection.get("target_boundary") or "", {})
                    target_name = target_product.get("name", "") or target_boundary.get(
                        "name", target_code
                    )
                    target_port = port_by_id.get(connection.get("target_port_id") or "", {})
                    expected_path = (
                        f"{connection.get('source_product', '')}({source_product.get('name', '')})/"
                        f"{source_port.get('name', '')} -> "
                        f"{target_code}({target_name})/"
                        f"{target_port.get('name', 'EXTERNAL_BOUNDARY' if connection.get('target_boundary') else '')} "
                        f"[{connection.get('item_code', '')}:"
                        f"{item_by_code[connection['item_code']].get('name_cn', '')}]"
                    )
                    expected = {
                        "Source_Leaf_Codes": connection.get("source_product", ""),
                        "Source_Leaf_Names": source_product.get("name", ""),
                        "Source_Raw_Port_IDs": connection.get("source_port_id", ""),
                        "Source_Raw_Port_Names": source_port.get("name", ""),
                        "Source_Raw_Directions": source_port.get("direction", ""),
                        "Target_Leaf_Codes": target_code,
                        "Target_Leaf_Names": target_name,
                        "Target_Raw_Port_IDs": connection.get("target_port_id") or "",
                        "Target_Raw_Port_Names": target_port.get("name", ""),
                        "Target_Raw_Directions": target_port.get("direction", ""),
                        "Original_Connection_Types": connection.get("mode", ""),
                        "Resolution_Basis": connection.get("resolution_basis", ""),
                        "Profile": connection.get("profile_id", ""),
                        "Original_Path_Summary": expected_path,
                    }
                    for column in paired_columns:
                        if actual[column][index] != str(expected[column]):
                            bad_fields.append(column)
                bad_fields = sorted(set(bad_fields))
            if bad_fields:
                pair_alignment_errors.append(
                    {
                        "aggregate_level": level_code,
                        "aggregate_interface_id": record["Aggregate_Interface_ID"],
                        "fields": sorted(set(bad_fields)),
                    }
                )

    pair_alignment_error = len(pair_alignment_errors)

    # Direction and attribute QA.
    attribute_failures = sum(
        record["Attribute_Propagation_Status"] != "PASS"
        for records in level_records.values()
        for record in records
    )
    unknown_direction = 0
    for records in level_records.values():
        for record in records:
            if record["Boundary_Type"] != "CROSS_BOUNDARY":
                continue
            if record["Physical_or_Signal"] == "Signal":
                valid = (
                    record["Source_Aggregate_Direction"] == "out"
                    and record["Target_Aggregate_Direction"] == "in"
                )
            else:
                valid = (
                    record["Source_Aggregate_Direction"] == "inout"
                    and record["Target_Aggregate_Direction"] == "inout"
                )
            if not valid:
                unknown_direction += 1

    all_aggregate_ids = [
        record["Aggregate_Interface_ID"]
        for records in level_records.values()
        for record in records
    ]
    duplicate_id_count = sum(count - 1 for count in Counter(all_aggregate_ids).values() if count > 1)

    # Cross-check current one-shot projections against existing upward inference.
    promoted_ports = upward["promoted_ports"]
    promoted_signal_index = {
        (row.get("derived_from_connection_id"), row.get("side"), row.get("owner_product_code"))
        for row in promoted_ports
    }
    upward_checks = 0
    upward_mismatches: list[dict[str, str]] = []
    for connection in final_connections:
        for level_code, level_number, _sheet_name in LEVELS:
            source_aggregate = ancestor(connection.get("source_product"), level_number)
            if connection.get("target_product"):
                target_aggregate = ancestor(connection.get("target_product"), level_number)
                target_code = (target_aggregate or {}).get("id", "")
            else:
                target_aggregate = None
                target_code = connection.get("target_boundary", "")
            source_code = (source_aggregate or {}).get("id", "")
            if not source_code or not target_code or source_code == target_code:
                continue

            if connection.get("mode") == "SIGNAL" and connection.get("target_product"):
                endpoints = [
                    ("SOURCE", source_code, connection.get("source_product"), connection.get("source_port_id")),
                    ("TARGET", target_code, connection.get("target_product"), connection.get("target_port_id")),
                ]
                for side, aggregate_code, leaf_code, _raw_port_id in endpoints:
                    if aggregate_code == leaf_code:
                        continue
                    upward_checks += 1
                    if (connection["connection_id"], side, aggregate_code) not in promoted_signal_index:
                        upward_mismatches.append(
                            {
                                "connection_id": connection["connection_id"],
                                "level": level_code,
                                "side": side,
                                "aggregate_code": aggregate_code,
                            }
                        )
            elif connection.get("mode") != "SIGNAL":
                endpoints = [
                    (source_code, connection.get("source_product"), connection.get("source_port_id"))
                ]
                if connection.get("target_product"):
                    endpoints.append(
                        (target_code, connection.get("target_product"), connection.get("target_port_id"))
                    )
                for aggregate_code, leaf_code, raw_port_id in endpoints:
                    if aggregate_code == leaf_code:
                        continue
                    upward_checks += 1
                    exists = any(
                        row.get("owner_product_code") == aggregate_code
                        and row.get("item_code") == connection.get("item_code")
                        and raw_port_id in row.get("source_leaf_port_ids", [])
                        for row in promoted_ports
                    )
                    if not exists:
                        upward_mismatches.append(
                            {
                                "connection_id": connection["connection_id"],
                                "level": level_code,
                                "side": "PHYSICAL",
                                "aggregate_code": aggregate_code,
                            }
                        )

    def projection(source: str, target: str) -> dict[str, str]:
        result: dict[str, str] = {}
        for level_code, level_number, _sheet_name in LEVELS:
            source_aggregate = ancestor(source, level_number)
            target_aggregate = ancestor(target, level_number)
            result[level_code] = (
                f"{(source_aggregate or {}).get('id', '')} → {(target_aggregate or {}).get('id', '')}"
            )
        return result

    anchor_4235_5111 = projection("4235", "5111")
    anchor_7211_8125 = projection("7211", "8125")
    expected_anchor_1 = {"L3": "4230 → 5110", "L2": "4200 → 5100", "L1": "4000 → 5000"}
    expected_anchor_2 = {"L3": "7210 → 8120", "L2": "7200 → 8100", "L1": "7000 → 8000"}
    anchor_1_connection = next(
        (
            row
            for row in final_connections
            if row.get("source_product") == "4235"
            and row.get("target_product") == "5111"
            and row.get("item_code") == "ITM-MEA-005"
        ),
        None,
    )
    anchor_2_connection = next(
        (
            row
            for row in final_connections
            if row.get("source_product") == "7211"
            and row.get("target_product") == "8125"
            and row.get("item_code") == "ITM-CMD-006"
        ),
        None,
    )
    false_test_case = sum(connection is None for connection in (anchor_1_connection, anchor_2_connection))

    stats: dict[str, dict[str, int]] = {}
    for level_code, _level_number, _sheet_name in LEVELS:
        records = level_records[level_code]
        stats[level_code] = {
            "total_aggregate_records": len(records),
            "cross_boundary": sum(row["Boundary_Type"] == "CROSS_BOUNDARY" for row in records),
            "internal": sum(row["Boundary_Type"] == "INTERNAL" for row in records),
            "unique_signal_interfaces": sum(row["Physical_or_Signal"] == "Signal" for row in records),
            "unique_physical_interfaces": sum(row["Physical_or_Signal"] == "Physical" for row in records),
        }

    qa_core_pass = all(
        [
            covered_all_levels == expected_connection_ids,
            not missing_item_ids,
            not missing_source_ancestor_ids,
            not missing_target_ancestor_ids,
            not unresolved_trace_ids,
            attribute_failures == 0,
            pair_alignment_error == 0,
            false_test_case == 0,
            duplicate_id_count == 0,
            unknown_direction == 0,
            not upward_mismatches,
            anchor_4235_5111 == expected_anchor_1,
            anchor_7211_8125 == expected_anchor_2,
            anchor_1_connection is not None,
            anchor_2_connection is not None,
        ]
    )

    qa = {
        "status": "PASS" if qa_core_pass else "FAIL",
        "original_final_connections_covered": len(covered_all_levels),
        "original_final_connections_expected": 184,
        "coverage_by_level": {
            level: len(coverage_by_level[level]) for level, _, _ in LEVELS
        },
        "missing_item": len(missing_item_ids),
        "missing_source_ancestor": len(missing_source_ancestor_ids),
        "missing_target_ancestor": len(missing_target_ancestor_ids),
        "unresolved_trace": len(unresolved_trace_ids),
        "attribute_propagation_failures": attribute_failures,
        "PAIR_ALIGNMENT_ERROR": pair_alignment_error,
        "pair_alignment_errors": pair_alignment_errors,
        "FALSE_TEST_CASE": false_test_case,
        "duplicate_aggregate_interface_id": duplicate_id_count,
        "unknown_direction": unknown_direction,
        "upward_inference_checks": upward_checks,
        "upward_inference_mismatches": upward_mismatches,
        "anchors": {
            "4235_to_5111": {
                "item": "高压线路电压",
                "present_in_frozen_final_connections": anchor_1_connection is not None,
                "original_connection_id": (anchor_1_connection or {}).get("connection_id", ""),
                "projection": anchor_4235_5111,
                "status": "PASS" if anchor_4235_5111 == expected_anchor_1 else "FAIL",
            },
            "7211_to_8125": {
                "item": "动态制动力请求",
                "present_in_frozen_final_connections": anchor_2_connection is not None,
                "original_connection_id": (anchor_2_connection or {}).get("connection_id", ""),
                "projection": anchor_7211_8125,
                "status": "PASS"
                if anchor_7211_8125 == expected_anchor_2 and anchor_2_connection is not None
                else "FAIL",
            },
        },
    }

    semantic_payload = {"levels": level_records, "statistics": stats, "qa": qa}
    semantic_digest = canonical_digest(semantic_payload)
    output_payload = {
        "meta": {
            "artifact_id": "Rail_MBSE_分层接口聚合_v1",
            "schema_version": "1.0.0",
            "source_counts": actual_counts,
            "source_sha256": {
                "architecture_xmi_ready_v1.json": sha256_file(ARCHITECTURE_PATH),
                "final_connections.json": sha256_file(FINAL_CONNECTIONS_PATH),
                "upward_port_inference_v1.json": sha256_file(UPWARD_PORTS_PATH),
                "hierarchical_connection_projection_v1.json": sha256_file(
                    HIERARCHICAL_PROJECTION_PATH
                ),
            },
            "aggregation_levels": [level for level, _, _ in LEVELS],
            "semantic_digest_sha256": semantic_digest,
        },
        **semantic_payload,
    }

    OUTPUT_JSON.parent.mkdir(parents=True, exist_ok=True)
    json_text = json.dumps(output_payload, ensure_ascii=False, indent=2, sort_keys=False) + "\n"
    json_temp = OUTPUT_JSON.with_suffix(".json.tmp")
    json_temp.write_text(json_text, encoding="utf-8", newline="\n")
    os.replace(json_temp, OUTPUT_JSON)

    write_workbook(level_records, OUTPUT_XLSX)
    verify_workbook(level_records, OUTPUT_XLSX)

    print_statistics(stats, qa, semantic_digest)
    return 0 if qa_core_pass else 2


def write_workbook(level_records: dict[str, list[dict[str, Any]]], output_path: Path) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.properties.creator = "Rail MBSE aggregation pipeline"
    workbook.properties.lastModifiedBy = "Rail MBSE aggregation pipeline"
    workbook.properties.created = FIXED_EXCEL_TIME
    workbook.properties.modified = FIXED_EXCEL_TIME
    workbook.properties.title = "Rail MBSE 分层接口聚合 v1"
    workbook.properties.subject = "L3/L2/L1 one-shot interface aggregation"

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    body_font = Font(name="Arial", size=10, color="1F1F1F")
    cross_fill = PatternFill("solid", fgColor="EAF3F8")
    internal_fill = PatternFill("solid", fgColor="F2F2F2")
    thin_blue = Side(style="thin", color="D9E2F3")
    header_side = Side(style="thin", color="FFFFFF")
    header_border = Border(left=header_side, right=header_side, bottom=header_side)
    body_border = Border(bottom=thin_blue)

    long_columns = {
        "Validity_Rule",
        "Fault_Tolerance_Strategy",
        "Semantic_Definition",
        "Aliases",
        "Usage_Note",
        "Original_Connection_IDs",
        "Source_Raw_Port_IDs",
        "Source_Raw_Port_Names",
        "Target_Raw_Port_IDs",
        "Target_Raw_Port_Names",
        "Original_Path_Summary",
        "Resolution_Basis",
        "Notes",
    }
    explicit_widths = {
        "Aggregate_Interface_ID": 20,
        "Aggregate_Level": 12,
        "Boundary_Type": 18,
        "Source_Aggregate_Code": 22,
        "Source_Aggregate_Name": 26,
        "Source_Aggregate_Port_Name": 30,
        "Source_Aggregate_Direction": 26,
        "Target_Aggregate_Code": 22,
        "Target_Aggregate_Name": 30,
        "Target_Aggregate_Port_Name": 30,
        "Target_Aggregate_Direction": 26,
        "Item_Code": 18,
        "Item_Name": 24,
        "Item_Name_EN": 30,
        "Item_Family": 18,
        "Physical_or_Signal": 20,
        "Port_Category": 18,
        "Domain": 18,
        "Datatype": 20,
        "Unit_or_Medium": 24,
        "Semantic_Definition": 52,
        "Aliases": 38,
        "Usage_Note": 44,
        "Port_Type_ID": 20,
        "Interface_Type_ID": 22,
        "Flow_Key": 44,
        "Leaf_Connection_Count": 22,
        "Original_Connection_IDs": 46,
        "Source_Leaf_Codes": 24,
        "Source_Leaf_Names": 30,
        "Source_Raw_Port_IDs": 50,
        "Source_Raw_Port_Names": 34,
        "Source_Raw_Directions": 24,
        "Target_Leaf_Codes": 28,
        "Target_Leaf_Names": 34,
        "Target_Raw_Port_IDs": 50,
        "Target_Raw_Port_Names": 34,
        "Target_Raw_Directions": 24,
        "Original_Connection_Types": 28,
        "Resolution_Basis": 48,
        "Original_Connection_Status": 28,
        "Profile": 24,
        "Active_Status": 18,
        "Original_Path_Summary": 70,
        "Validity_Rule": 44,
        "Fault_Tolerance_Strategy": 34,
        "Notes": 38,
    }

    for level_code, _level_number, sheet_name in LEVELS:
        sheet = workbook.create_sheet(sheet_name)
        sheet.sheet_view.showGridLines = False
        sheet.freeze_panes = "A2"
        sheet.sheet_properties.tabColor = "1F4E78" if level_code == "L3" else "9DC3E6"
        sheet.append(COLUMNS)
        for record in level_records[level_code]:
            sheet.append([record.get(column, "") for column in COLUMNS])

        max_row = sheet.max_row
        max_col = len(COLUMNS)
        sheet.auto_filter.ref = f"A1:{get_column_letter(max_col)}{max_row}"
        sheet.row_dimensions[1].height = 42

        for cell in sheet[1]:
            cell.fill = copy(header_fill)
            cell.font = copy(header_font)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = copy(header_border)

        boundary_col = COLUMNS.index("Boundary_Type") + 1
        for row_index in range(2, max_row + 1):
            is_internal = sheet.cell(row_index, boundary_col).value == "INTERNAL"
            row_fill = internal_fill if is_internal else cross_fill
            sheet.row_dimensions[row_index].height = 30
            for col_index in range(1, max_col + 1):
                cell = sheet.cell(row_index, col_index)
                cell.font = copy(body_font)
                cell.fill = copy(row_fill)
                cell.border = copy(body_border)
                cell.alignment = Alignment(
                    horizontal="left",
                    vertical="top" if COLUMNS[col_index - 1] in long_columns else "center",
                    wrap_text=COLUMNS[col_index - 1] in long_columns,
                )

        for col_index, column in enumerate(COLUMNS, 1):
            if column in explicit_widths:
                width = explicit_widths[column]
            else:
                sample_values = [column] + [
                    str(record.get(column, "")) for record in level_records[level_code][:120]
                ]
                width = min(32, max(12, max(len(value) for value in sample_values) * 1.05))
            sheet.column_dimensions[get_column_letter(col_index)].width = width

        if max_row >= 2:
            data_range = f"A2:{get_column_letter(max_col)}{max_row}"
            # Rules remain dynamic if reviewers sort/filter or edit Boundary_Type.
            sheet.conditional_formatting.add(
                data_range,
                FormulaRule(formula=["$C2=\"INTERNAL\""], fill=internal_fill),
            )
            sheet.conditional_formatting.add(
                data_range,
                FormulaRule(formula=["$C2=\"CROSS_BOUNDARY\""], fill=cross_fill),
            )

        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.auto_filter.ref = f"A1:{get_column_letter(max_col)}{max_row}"

    output_path.parent.mkdir(parents=True, exist_ok=True)
    temp_path = output_path.with_suffix(".xlsx.tmp")
    workbook.save(temp_path)
    os.replace(temp_path, output_path)


def verify_workbook(level_records: dict[str, list[dict[str, Any]]], output_path: Path) -> None:
    workbook = load_workbook(output_path, read_only=False, data_only=False)
    expected_sheets = [sheet_name for _, _, sheet_name in LEVELS]
    if workbook.sheetnames != expected_sheets:
        raise RuntimeError(f"Workbook sheet mismatch: {workbook.sheetnames}")
    for level_code, _level_number, sheet_name in LEVELS:
        sheet = workbook[sheet_name]
        headers = [cell.value for cell in sheet[1]]
        if headers != COLUMNS:
            raise RuntimeError(f"Column mismatch in {sheet_name}")
        if sheet.max_row - 1 != len(level_records[level_code]):
            raise RuntimeError(f"Row-count mismatch in {sheet_name}")
        if sheet.freeze_panes != "A2" or not sheet.auto_filter.ref:
            raise RuntimeError(f"Review formatting missing in {sheet_name}")
        id_col = COLUMNS.index("Aggregate_Interface_ID") + 1
        ids = [sheet.cell(row, id_col).value for row in range(2, sheet.max_row + 1)]
        if len(ids) != len(set(ids)):
            raise RuntimeError(f"Duplicate Aggregate_Interface_ID in {sheet_name}")
    workbook.close()


def print_statistics(
    stats: dict[str, dict[str, int]], qa: dict[str, Any], semantic_digest: str
) -> None:
    for level_code, _level_number, _sheet_name in LEVELS:
        row = stats[level_code]
        print(f"{level_code}:")
        print(f"Total aggregate records: {row['total_aggregate_records']}")
        print(f"Cross-boundary: {row['cross_boundary']}")
        print(f"Internal: {row['internal']}")
        print(f"Unique signal interfaces: {row['unique_signal_interfaces']}")
        print(f"Unique physical interfaces: {row['unique_physical_interfaces']}")
    print(
        "Original Final Connections covered: "
        f"{qa['original_final_connections_covered']}/{qa['original_final_connections_expected']}"
    )
    print(f"Missing Item: {qa['missing_item']}")
    print(f"Missing Source Ancestor: {qa['missing_source_ancestor']}")
    print(f"Missing Target Ancestor: {qa['missing_target_ancestor']}")
    print(f"Attribute propagation failures: {qa['attribute_propagation_failures']}")
    print(f"PAIR_ALIGNMENT_ERROR: {qa['PAIR_ALIGNMENT_ERROR']}")
    print(f"FALSE_TEST_CASE: {qa['FALSE_TEST_CASE']}")
    print(f"Duplicate Aggregate_Interface_ID: {qa['duplicate_aggregate_interface_id']}")
    print(f"Unknown direction: {qa['unknown_direction']}")
    print(f"JSON semantic digest: {semantic_digest}")
    print(f"QA: {qa['status']}")


if __name__ == "__main__":
    raise SystemExit(main())
