"""Run all deterministic Function Architecture stages twice and verify repeatability."""
from __future__ import annotations

import hashlib
import json
import subprocess
import sys
from pathlib import Path

from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
ARCH = ROOT / "work" / "architecture_xmi_ready.json"
SCHEMA = ROOT / "work" / "architecture_xmi_ready.schema.json"
BOOK = ROOT / "Rail_MBSE_Function_Interface_v1.xlsx"
QA = HERE / "work" / "function_architecture_qa.json"
REPORT = ROOT / "reports" / "function_architecture_report.md"


def sha256(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def workbook_semantic_sha(path):
    workbook = load_workbook(path, read_only=False, data_only=False)
    payload = []
    for sheet in workbook.worksheets:
        payload.append([sheet.title, [list(row) for row in sheet.iter_rows(values_only=True)]])
    workbook.close()
    raw = json.dumps(payload, ensure_ascii=False, separators=(",", ":"), default=str).encode("utf-8")
    return hashlib.sha256(raw).hexdigest()


def execute_pass(pass_number):
    for number in range(1, 7):
        script = next(HERE.glob(f"{number:02d}_*.py"))
        subprocess.run([sys.executable, str(script)], cwd=ROOT, check=True)
    builder = (HERE / "07_export_interface_workbook.py").read_text(encoding="utf-8")
    subprocess.run(["node", "--input-type=module", "--eval", builder], cwd=HERE, check=True)
    subprocess.run([sys.executable, str(HERE / "08_validate_function_architecture.py")], cwd=ROOT, check=True)
    result = {
        "architecture_json_sha256": sha256(ARCH),
        "schema_json_sha256": sha256(SCHEMA),
        "workbook_sha256": sha256(BOOK),
        "workbook_semantic_sha256": workbook_semantic_sha(BOOK),
    }
    print(json.dumps({"pass": pass_number, "digests": result}, ensure_ascii=False))
    return result


def main():
    first = execute_pass(1)
    second = execute_pass(2)
    repeatability = {
        "architecture_json_byte_identical": first["architecture_json_sha256"] == second["architecture_json_sha256"],
        "schema_json_byte_identical": first["schema_json_sha256"] == second["schema_json_sha256"],
        "workbook_byte_identical": first["workbook_sha256"] == second["workbook_sha256"],
        "workbook_semantically_identical": first["workbook_semantic_sha256"] == second["workbook_semantic_sha256"],
    }
    repeatability["status"] = "PASS" if repeatability["architecture_json_byte_identical"] and repeatability["schema_json_byte_identical"] and repeatability["workbook_semantically_identical"] else "FAIL"
    qa = json.loads(QA.read_text(encoding="utf-8"))
    qa["repeatability"] = repeatability
    qa["xmi_qa"]["XMI-QA-09"] = repeatability["status"] == "PASS"
    qa["status"] = "PASS" if all(qa["function_qa"].values()) and all(qa["xmi_qa"].values()) and qa["workbook_qa"] == "PASS" else "FAIL"
    QA.write_text(json.dumps(qa, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    report = REPORT.read_text(encoding="utf-8").rstrip() + f"\n- Two-pass repeatability: {repeatability['status']} (JSON byte-identical; workbook semantic-identical)\n"
    REPORT.write_text(report, encoding="utf-8")
    inspect_sidecar = ROOT / "Rail_MBSE_Function_Interface_v1.xlsx.inspect.ndjson"
    if inspect_sidecar.exists():
        inspect_sidecar.unlink()
    print(json.dumps({"pipeline_status": qa["status"], "repeatability": repeatability, "counts": qa["counts"]}, ensure_ascii=False))
    if qa["status"] != "PASS":
        raise SystemExit(1)


if __name__ == "__main__":
    main()
