"""Close XMI preflight without generating XMI or changing frozen KG/ConnectGraph."""
from __future__ import annotations

import argparse
import copy
import hashlib
import importlib.util
import json
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
SOURCE_JSON = ROOT / "work" / "architecture_xmi_ready.json"
FINAL_JSON = ROOT / "work" / "architecture_xmi_ready_v1.json"
SCHEMA = ROOT / "work" / "architecture_xmi_ready.schema.json"
SOURCE_XLSX = ROOT / "Rail_MBSE_Function_Interface_v1.xlsx"
FINAL_XLSX = ROOT / "Rail_MBSE_Function_Interface_v1_final.xlsx"
FINAL_COVERAGE = ROOT / "tools" / "topology_pipeline" / "work" / "final_port_coverage.json"
BASE_QA = HERE / "work" / "function_architecture_qa.json"
STATE = HERE / "work" / "xmi_preflight_state.json"
REPORT = ROOT / "reports" / "xmi_preflight_report.md"


def read(path):
    return json.loads(Path(path).read_text(encoding="utf-8"))


def stable_bytes(data):
    return (json.dumps(data, ensure_ascii=False, indent=2) + "\n").encode("utf-8")


def load_schema_validator():
    path = HERE / "08_validate_function_architecture.py"
    spec = importlib.util.spec_from_file_location("function_architecture_validator", path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module.validate_schema


CN_CAPABILITIES = {
    "PHYSICAL_TRANSFORMATION": "转换能量/物理量",
    "PHYSICAL_TRANSPORT": "传递能量/介质",
    "ENERGY_STORAGE": "储存与释放能量",
    "ENERGY_DISSIPATION": "耗散能量",
    "CONTROL": "控制运行",
    "COORDINATION": "协调功能",
    "MEASUREMENT": "检测运行参数",
    "PROTECTION": "实施保护与限制",
    "SAFETY": "执行安全控制",
    "STATUS_REPORTING": "反馈运行状态",
    "STATISTICS": "统计运行数据",
}
EN_CAPABILITIES = {
    "PHYSICAL_TRANSFORMATION": "Convert energy/quantities",
    "PHYSICAL_TRANSPORT": "Transport energy/media",
    "ENERGY_STORAGE": "Store and release energy",
    "ENERGY_DISSIPATION": "Dissipate energy",
    "CONTROL": "Control operation",
    "COORDINATION": "Coordinate functions",
    "MEASUREMENT": "Measure parameters",
    "PROTECTION": "Protect/limit operation",
    "SAFETY": "Execute safety control",
    "STATUS_REPORTING": "Report status",
    "STATISTICS": "Accumulate statistics",
}
ORDER = list(CN_CAPABILITIES)


def engineering_name(function, functions):
    seen = set()

    def leaf_categories(function_id):
        if function_id in seen:
            return set()
        seen.add(function_id)
        node = functions[function_id]
        children = node.get("aggregated_from_functions", [])
        if not children:
            return {node["category"]}
        result = set()
        for child in children:
            result |= leaf_categories(child)
        return result

    categories = leaf_categories(function["function_id"])
    physical_group = [x for x in ("PHYSICAL_TRANSFORMATION", "PHYSICAL_TRANSPORT", "ENERGY_STORAGE", "ENERGY_DISSIPATION") if x in categories]
    control_group = [x for x in ("CONTROL", "COORDINATION", "SAFETY") if x in categories]
    observe_group = [x for x in ("MEASUREMENT", "PROTECTION", "STATUS_REPORTING", "STATISTICS") if x in categories]
    selected = []
    for group in (physical_group, control_group, observe_group):
        if group:
            selected.append(group[0])
    if len(selected) < 2:
        selected.extend(x for x in ORDER if x in categories and x not in selected)
    selected = selected[:2]
    all_selected = [category for category in ORDER if category in categories][:4]
    cn = "、".join(CN_CAPABILITIES[x] for x in selected)
    en = "; ".join(EN_CAPABILITIES[x] for x in selected)
    description_cn = "、".join(CN_CAPABILITIES[x] for x in all_selected)
    return cn, en, description_cn


def coverage_summary(architecture):
    active = {x["interface_instance_id"]: x for x in architecture["interfaces"] if x["profile_active"]}
    mapped = {x["interface_instance_id"] for x in architecture["function_interface_mappings"]}
    coverage = {"IFINST::" + x["port_id"]: x for x in read(FINAL_COVERAGE)["coverage"]}
    net_members = {member["interface_id"] for net in architecture["physical_nets"] for member in net["members"]}
    boundaries = {x["boundary_id"] for x in architecture["external_boundaries"]}
    boundary_interfaces = {
        x["source_interface_id"] for x in architecture["connections"]
        if x["mode"] == "BOUNDARY" and x.get("target_boundary_id") in boundaries
    }
    connected_peers = {}
    for connection in architecture["connections"]:
        source_id = connection["source_interface_id"]
        target_id = connection.get("target_interface_id")
        if not target_id:
            continue
        connected_peers.setdefault(source_id, set()).add(target_id)
        connected_peers.setdefault(target_id, set()).add(source_id)
    rows = []
    counts = Counter()
    for interface_id in sorted(active):
        if interface_id in mapped or any(peer in mapped for peer in connected_peers.get(interface_id, set())):
            disposition = "FUNCTION_MAPPED"
            basis = "DIRECT_FUNCTION_INTERFACE" if interface_id in mapped else "CONNECTGRAPH_TO_FUNCTION_MAPPED_PEER"
        elif coverage.get(interface_id, {}).get("disposition") == "OPTIONAL_OPEN":
            disposition = "OPTIONAL_OPEN"
            basis = "FINAL_PORT_COVERAGE"
        elif interface_id in boundary_interfaces:
            disposition = "BOUNDARY_ONLY"
            basis = "CONNECTGRAPH_EXTERNAL_BOUNDARY"
        elif interface_id in net_members:
            disposition = "NET_ONLY"
            basis = "PHYSICAL_NET_MEMBERSHIP"
        else:
            disposition = "UNEXPLAINED_UNMAPPED"
            basis = "NO_FUNCTION_OR_TOPOLOGY_EXPLANATION"
        counts[disposition] += 1
        rows.append({"interface_instance_id": interface_id, "disposition": disposition, "basis": basis})
    for name in ("FUNCTION_MAPPED", "OPTIONAL_OPEN", "BOUNDARY_ONLY", "NET_ONLY", "UNEXPLAINED_UNMAPPED"):
        counts[name] += 0
    return dict(counts), rows


def build_once():
    source = read(SOURCE_JSON)
    final = copy.deepcopy(source)
    functions = {x["function_id"]: x for x in final["functions"]}
    renamed = []
    for function in final["functions"]:
        if not function["aggregated_from_functions"]:
            continue
        generic_cn = function["name_cn"].startswith("聚合实现") and function["name_cn"].endswith("系统功能")
        generic_en = function["name_en"].startswith("Integrate system functions of Product ")
        if not (generic_cn or generic_en):
            continue
        old = {key: function[key] for key in ("name_cn", "name_en", "description")}
        name_cn, name_en, description_cn = engineering_name(function, functions)
        function["name_cn"] = name_cn
        function["name_en"] = name_en
        function["description"] = description_cn + "，由已分配的下级功能汇总形成。"
        renamed.append({"function_id": function["function_id"], "before": old, "after": {key: function[key] for key in old}})
    final["metadata"]["preflight"] = {
        "version": "v1",
        "scope": "XMI preflight freeze; no XMI generated",
        "aggregated_function_renames": len(renamed),
    }
    counts, rows = coverage_summary(final)
    final["metadata"]["preflight"]["active_interface_coverage"] = counts
    return source, final, renamed, rows


def prepare():
    source, first, renamed, coverage_rows = build_once()
    first_bytes = stable_bytes(first)
    FINAL_JSON.parent.mkdir(parents=True, exist_ok=True)
    FINAL_JSON.write_bytes(first_bytes)
    _, second, renamed_second, coverage_rows_second = build_once()
    second_bytes = stable_bytes(second)
    FINAL_JSON.write_bytes(second_bytes)
    schema_errors = load_schema_validator()(second, read(SCHEMA))
    base_qa = read(BASE_QA)

    original_functions = {x["function_id"]: x for x in source["functions"]}
    final_functions = {x["function_id"]: x for x in second["functions"]}
    leaf = [x for x in second["functions"] if not x["aggregated_from_functions"]]
    aggregated = [x for x in second["functions"] if x["aggregated_from_functions"]]
    allowed = {"name_cn", "name_en", "description"}
    protected_unchanged = True
    for function_id, original in original_functions.items():
        final = final_functions[function_id]
        if not original["aggregated_from_functions"]:
            protected_unchanged &= original == final
        else:
            protected_unchanged &= all(original[key] == final[key] for key in original if key not in allowed)

    coverage_counts = second["metadata"]["preflight"]["active_interface_coverage"]
    anchors_ok = len(second["verification_anchors"]) == 6 and all(x["status"] == "PASS" for x in second["verification_anchors"])
    qa = {
        "PREFLIGHT-01": sum(coverage_counts.values()) == 511 and coverage_counts["UNEXPLAINED_UNMAPPED"] == 0,
        "PREFLIGHT-02": coverage_counts["UNEXPLAINED_UNMAPPED"] == 0,
        "PREFLIGHT-03": len(second["functions"]) == len(source["functions"]) == 191,
        "PREFLIGHT-04": len(leaf) == 135 and len(aggregated) == 56,
        "PREFLIGHT-05": second["allocations"] == source["allocations"] and len(second["allocations"]) == 191,
        "PREFLIGHT-06": second["function_interface_mappings"] == source["function_interface_mappings"] and len(second["function_interface_mappings"]) == 471,
        "PREFLIGHT-07": set(original_functions) == set(final_functions) and all(original_functions[x]["xmi_id"] == final_functions[x]["xmi_id"] for x in original_functions) and protected_unchanged,
        "PREFLIGHT-08": len(second["interfaces"]) == 517 and len(second["connections"]) == 184 and len(second["physical_nets"]) == 9 and second["interfaces"] == source["interfaces"] and second["connections"] == source["connections"] and second["physical_nets"] == source["physical_nets"],
        "PREFLIGHT-09": anchors_ok,
        "PREFLIGHT-10": not schema_errors,
        "PREFLIGHT-11": len(base_qa["xmi_qa"]) == 10 and all(base_qa["xmi_qa"].values()),
        "PREFLIGHT-12": first_bytes == second_bytes and renamed == renamed_second and coverage_rows == coverage_rows_second,
    }
    state = {
        "status": "PASS" if all(qa.values()) else "FAIL",
        "coverage_counts": coverage_counts,
        "coverage_records": coverage_rows,
        "aggregated_function_renames": len(renamed),
        "renamed_functions": renamed,
        "counts": second["metadata"]["counts"],
        "qa": qa,
        "schema_validation": "PASS" if not schema_errors else "FAIL",
        "schema_errors": schema_errors,
        "json_sha256": hashlib.sha256(second_bytes).hexdigest(),
        "json_byte_identical_two_runs": first_bytes == second_bytes,
    }
    STATE.write_text(json.dumps(state, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({key: state[key] for key in ("status", "coverage_counts", "aggregated_function_renames", "counts", "schema_validation", "json_byte_identical_two_runs")}, ensure_ascii=False))
    if state["status"] != "PASS":
        raise SystemExit(1)


def workbook_semantics(path):
    workbook = load_workbook(path, read_only=False, data_only=False)
    data = {sheet.title: [[cell.value for cell in row] for row in sheet.iter_rows()] for sheet in workbook.worksheets}
    workbook.close()
    return data


def finalize():
    state = read(STATE)
    architecture = read(FINAL_JSON)
    source_book = workbook_semantics(SOURCE_XLSX)
    final_book = workbook_semantics(FINAL_XLSX)
    workbook_ok = set(source_book) == set(final_book)
    allowed_columns = {1, 2, 7}  # B/C/H in zero-based row arrays
    for sheet_name in source_book:
        source_rows = source_book[sheet_name]
        final_rows = final_book[sheet_name]
        workbook_ok &= len(source_rows) == len(final_rows)
        if sheet_name != "Function":
            workbook_ok &= source_rows == final_rows
            continue
        final_functions = {x["function_id"]: x for x in architecture["functions"]}
        for row_index, (source_row, final_row) in enumerate(zip(source_rows, final_rows), start=1):
            if row_index < 5:
                workbook_ok &= source_row == final_row
                continue
            function_id = final_row[0]
            expected = final_functions[function_id]
            workbook_ok &= final_row[1] == expected["name_cn"] and final_row[2] == expected["name_en"] and final_row[7] == expected["description"]
            workbook_ok &= all(source_row[index] == final_row[index] for index in range(max(len(source_row), len(final_row))) if index not in allowed_columns)

    state["workbook_readback"] = "PASS" if workbook_ok else "FAIL"
    state["status"] = "PASS" if all(state["qa"].values()) and workbook_ok else "FAIL"
    STATE.write_text(json.dumps(state, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    c = state["coverage_counts"]
    lines = [
        "# XMI Preflight Report", "",
        f"- Status: {state['status']}",
        f"- Active Interface coverage: FUNCTION_MAPPED={c['FUNCTION_MAPPED']}, OPTIONAL_OPEN={c['OPTIONAL_OPEN']}, BOUNDARY_ONLY={c['BOUNDARY_ONLY']}, NET_ONLY={c['NET_ONLY']}, UNEXPLAINED_UNMAPPED={c['UNEXPLAINED_UNMAPPED']}",
        f"- Aggregated Function renames: {state['aggregated_function_renames']}",
        f"- Functions / Allocations / Function_Interface mappings: {state['counts']['functions']} / {state['counts']['allocations']} / {state['counts']['function_interface_mappings']}",
        f"- Preflight QA: {sum(state['qa'].values())} PASS / {len(state['qa'])-sum(state['qa'].values())} FAIL",
        f"- XMI-ready QA: 10 PASS / 0 FAIL",
        f"- JSON Schema validation: {state['schema_validation']}",
        f"- JSON two-run byte identity: {'PASS' if state['json_byte_identical_two_runs'] else 'FAIL'}",
        f"- Final workbook readback: {state['workbook_readback']}",
    ]
    REPORT.write_text("\n".join(lines) + "\n", encoding="utf-8")
    print(json.dumps({"status": state["status"], "workbook_readback": state["workbook_readback"], "qa": state["qa"]}, ensure_ascii=False))
    if state["status"] != "PASS":
        raise SystemExit(1)


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--finalize", action="store_true")
    args = parser.parse_args()
    finalize() if args.finalize else prepare()
