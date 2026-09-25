"""Regression checks for the generated interface-layer workbook."""

from __future__ import annotations

import json
import sys
import unittest
from pathlib import Path


PIPELINE_DIR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PIPELINE_DIR))

from common import OUTPUT_WORKBOOK, RAW_PORT_COLUMNS, RAW_PORT_SHEET, WORK_DIR  # noqa: E402


class InterfacePipelineTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        from openpyxl import load_workbook

        if not OUTPUT_WORKBOOK.exists():
            raise AssertionError(f"先运行pipeline生成: {OUTPUT_WORKBOOK}")
        cls.summary = json.loads((WORK_DIR / "pipeline_summary.json").read_text(encoding="utf-8"))
        cls.wb = load_workbook(OUTPUT_WORKBOOK, read_only=True, data_only=False)

    @classmethod
    def tearDownClass(cls) -> None:
        cls.wb.close()

    def test_required_sheets_exist(self) -> None:
        required = {
            RAW_PORT_SHEET,
            "Architecture_Profile",
            "PortType库",
            "InterfaceType库",
            "Compatibility_Rules",
            "PortType映射",
            "Interface_QA",
        }
        self.assertTrue(required.issubset(self.wb.sheetnames))

    def test_no_connection_sheet(self) -> None:
        self.assertFalse([name for name in self.wb.sheetnames if "Connection" in name or "连接" in name])

    def test_raw_port_shape_and_headers(self) -> None:
        ws = self.wb[RAW_PORT_SHEET]
        headers = [cell.value for cell in ws[1]]
        self.assertEqual(headers, RAW_PORT_COLUMNS)
        self.assertEqual(ws.max_row - 1, self.summary["after_port_count"])
        self.assertEqual(ws.max_column, 9)

    def test_mapping_cardinality(self) -> None:
        ws = self.wb["PortType映射"]
        self.assertEqual(ws.max_row - 1, self.summary["after_port_count"])

    def test_stable_type_ids(self) -> None:
        ws = self.wb["PortType库"]
        ids = [row[0] for row in ws.iter_rows(min_row=2, values_only=True)]
        self.assertEqual(len(ids), len(set(ids)))
        self.assertTrue(all(value.startswith("PT-") for value in ids))
        self.assertEqual(len(ids), self.summary["port_type_count"])
        if_ws = self.wb["InterfaceType库"]
        if_ids = [row[0] for row in if_ws.iter_rows(min_row=2, values_only=True)]
        self.assertEqual(len(if_ids), len(set(if_ids)))
        self.assertTrue(all(value.startswith("IF-") for value in if_ids))
        self.assertEqual(len(if_ids), self.summary["interface_type_count"])

    def test_profile_has_no_active_dc_supply(self) -> None:
        ws = self.wb["PortType映射"]
        rows = list(ws.iter_rows(min_row=2, values_only=True))
        active_dc = [row for row in rows if row[6] == "直流牵引供电电能" and row[9] is True]
        inactive = [row for row in rows if row[9] is False]
        self.assertFalse(active_dc)
        self.assertEqual(len(inactive), self.summary["inactive_variant_port_count"])

    def test_required_qa_passes(self) -> None:
        ws = self.wb["Interface_QA"]
        statuses = {row[0]: row[1] for row in ws.iter_rows(min_row=2, values_only=True) if str(row[0]).startswith("QA-")}
        self.assertEqual(set(statuses), {f"QA-{index:02d}" for index in range(1, 15)})
        self.assertTrue(all(status == "PASS" for status in statuses.values()))


if __name__ == "__main__":
    unittest.main()
