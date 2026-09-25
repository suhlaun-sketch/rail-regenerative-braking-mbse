"""Inspect source workbooks and persist a reproducible input inventory."""

from __future__ import annotations

import argparse
from collections import Counter
from pathlib import Path
from typing import Any

from common import (
    AC_DC_ITEMS,
    AUXILIARY_WORKBOOKS,
    ITEM_WORKBOOK,
    PRIMARY_WORKBOOK,
    RAW_PORT_COLUMNS,
    RAW_PORT_SHEET,
    TARGET_REVIEW_CODES,
    WORK_DIR,
    ensure_required_inputs,
    normalize_code,
    normalize_text,
    write_json,
)


def row_dict(headers: list[str], values: tuple[Any, ...]) -> dict[str, str]:
    return {header: normalize_text(values[index]) if index < len(values) else "" for index, header in enumerate(headers)}


def inspect_workbook(path: Path) -> dict[str, Any]:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=False, data_only=False)
    result: dict[str, Any] = {"path": str(path), "sheets": []}
    for ws in wb.worksheets:
        max_row = ws.max_row or 0
        max_column = ws.max_column or 0
        first_rows = (
            list(ws.iter_rows(min_row=1, max_row=min(max_row, 5), values_only=True))
            if max_row and max_column
            else []
        )
        result["sheets"].append(
            {
                "name": ws.title,
                "max_row": max_row,
                "max_column": max_column,
                "first_rows": [[normalize_text(v) for v in row] for row in first_rows],
            }
        )
    wb.close()
    return result


def audit_item_dictionary() -> dict[str, Any]:
    from openpyxl import load_workbook

    wb = load_workbook(ITEM_WORKBOOK, read_only=True, data_only=False)
    if "Item字典" not in wb.sheetnames:
        raise ValueError("Item字典工作簿缺少 Item字典 Sheet")
    ws = wb["Item字典"]
    headers = [normalize_text(cell.value) for cell in ws[1]]
    required = ["item_code", "Item中文名", "领域", "数据类型", "单位/介质", "规范语义定义"]
    missing = [column for column in required if column not in headers]
    if missing:
        raise ValueError(f"Item字典缺少字段: {missing}")
    items = []
    for values in ws.iter_rows(min_row=2, values_only=True):
        record = row_dict(headers, values)
        if record["item_code"] or record["Item中文名"]:
            items.append(record)
    wb.close()
    names = [item["Item中文名"] for item in items]
    codes = [item["item_code"] for item in items]
    duplicate_names = sorted(name for name, count in Counter(names).items() if count > 1)
    duplicate_codes = sorted(code for code, count in Counter(codes).items() if count > 1)
    return {
        "item_count": len(items),
        "duplicate_item_names": duplicate_names,
        "duplicate_item_codes": duplicate_codes,
        "items": items,
    }


def selected_leaf_nodes() -> dict[str, Any]:
    from openpyxl import load_workbook

    wb = load_workbook(PRIMARY_WORKBOOK, read_only=True, data_only=True)
    if "入选" not in wb.sheetnames:
        raise ValueError("主输入缺少 入选 Sheet")
    ws = wb["入选"]
    headers = [normalize_text(cell.value) for cell in ws[1]]
    required = ["code", "name", "parent_name", "节点角色"]
    if any(column not in headers for column in required):
        raise ValueError(f"入选 Sheet字段不完整: {headers}")
    nodes = []
    for values in ws.iter_rows(min_row=2, values_only=True):
        record = row_dict(headers, values)
        record["code"] = normalize_code(record["code"])
        if record["code"]:
            nodes.append(record)
    parent_names = {node["parent_name"] for node in nodes if node["parent_name"]}
    leaves = [node for node in nodes if node["name"] not in parent_names]
    wb.close()
    return {"selected_node_count": len(nodes), "leaf_node_count": len(leaves), "nodes": nodes, "leaves": leaves}


def audit_primary() -> dict[str, Any]:
    from openpyxl import load_workbook

    wb = load_workbook(PRIMARY_WORKBOOK, read_only=True, data_only=False)
    if RAW_PORT_SHEET not in wb.sheetnames:
        raise ValueError(f"主输入缺少 {RAW_PORT_SHEET}")
    ws = wb[RAW_PORT_SHEET]
    headers = [normalize_text(cell.value) for cell in ws[1]]
    if headers != RAW_PORT_COLUMNS:
        raise ValueError(f"{RAW_PORT_SHEET} 9列定义不匹配: {headers}")

    ports: list[dict[str, str]] = []
    for values in ws.iter_rows(min_row=2, values_only=True):
        if not any(normalize_text(value) for value in values):
            continue
        record = row_dict(headers, values)
        record["code"] = normalize_code(record["code"])
        ports.append(record)

    code_counts = Counter(port["code"] for port in ports)
    item_counts = Counter(port["交换内容"] for port in ports)
    direction_counts = Counter(port["方向"] for port in ports)
    category_counts = Counter(port["端口类别"] for port in ports)
    duplicates = [
        {"key": list(key), "count": count}
        for key, count in Counter(tuple(port[column] for column in RAW_PORT_COLUMNS) for port in ports).items()
        if count > 1
    ]
    targeted = {code: [port for port in ports if port["code"] == code] for code in TARGET_REVIEW_CODES}
    ac_dc = [port for port in ports if port["交换内容"] in AC_DC_ITEMS]

    possible_leaf_codes: set[str] = set()
    for sheet in wb.worksheets:
        if sheet.title == RAW_PORT_SHEET:
            continue
        header_values = [normalize_text(cell.value) for cell in sheet[1]]
        if "code" in header_values:
            code_index = header_values.index("code")
            for values in sheet.iter_rows(min_row=2, values_only=True):
                if code_index < len(values):
                    code = normalize_code(values[code_index])
                    if code:
                        possible_leaf_codes.add(code)
    qa_rows: list[list[str]] = []
    if "RawPort_QA" in wb.sheetnames:
        qa_ws = wb["RawPort_QA"]
        qa_rows = [
            [normalize_text(value) for value in values]
            for values in qa_ws.iter_rows(values_only=True)
            if any(normalize_text(value) for value in values)
        ]
    wb.close()

    return {
        "raw_port_count": len(ports),
        "unique_port_nodes": len(code_counts),
        "possible_leaf_code_count_from_other_sheets": len(possible_leaf_codes),
        "ports_by_code": dict(sorted(code_counts.items())),
        "items_used": dict(sorted(item_counts.items())),
        "directions": dict(direction_counts),
        "categories": dict(category_counts),
        "exact_duplicates": duplicates,
        "targeted_review": targeted,
        "ac_dc_ports": ac_dc,
        "source_rawport_qa": qa_rows,
        "ports": ports,
    }


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, default=WORK_DIR / "input_inventory.json")
    args = parser.parse_args()

    ensure_required_inputs()
    inventories = [inspect_workbook(PRIMARY_WORKBOOK), inspect_workbook(ITEM_WORKBOOK)]
    inventories.extend(inspect_workbook(path) for path in AUXILIARY_WORKBOOKS if path.exists())
    audit = audit_primary()
    item_audit = audit_item_dictionary()
    selection = selected_leaf_nodes()
    item_by_name = {item["Item中文名"]: item for item in item_audit["items"]}
    audit["unmapped_item_ports"] = [port for port in audit["ports"] if port["交换内容"] not in item_by_name]
    audit["unit_medium_mismatches"] = [
        {
            "code": port["code"],
            "name": port["name"],
            "port_name": port["端口名称"],
            "item": port["交换内容"],
            "raw_unit_medium": port["单位/介质"],
            "canonical_unit_medium": item_by_name[port["交换内容"]]["单位/介质"],
        }
        for port in audit["ports"]
        if port["交换内容"] in item_by_name
        and port["单位/介质"] != item_by_name[port["交换内容"]]["单位/介质"]
    ]
    leaf_codes = {node["code"] for node in selection["leaves"]}
    audit["non_leaf_ports"] = [port for port in audit["ports"] if port["code"] not in leaf_codes]
    payload = {
        "workbooks": inventories,
        "item_dictionary": item_audit,
        "selection": selection,
        "primary_audit": audit,
    }
    write_json(args.output, payload)

    print(f"输入工作簿数: {len(inventories)}")
    for workbook in inventories:
        print(f"- {workbook['path']}")
        for sheet in workbook["sheets"]:
            print(f"  {sheet['name']}: {sheet['max_row']}行 x {sheet['max_column']}列")
    print(f"Raw Port数: {audit['raw_port_count']}")
    print(f"Raw Port节点数: {audit['unique_port_nodes']}")
    print(f"入选节点数: {selection['selected_node_count']}")
    print(f"入选叶节点数: {selection['leaf_node_count']}")
    print(f"Canonical Item数: {item_audit['item_count']}")
    print(f"实际使用Item数: {len(audit['items_used'])}")
    print(f"完全重复Port组数: {len(audit['exact_duplicates'])}")
    print(f"未映射Item Port数: {len(audit['unmapped_item_ports'])}")
    print(f"单位/介质不一致Port数: {len(audit['unit_medium_mismatches'])}")
    print(f"非叶节点Port数: {len(audit['non_leaf_ports'])}")
    print(f"审计清单: {args.output.resolve()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
