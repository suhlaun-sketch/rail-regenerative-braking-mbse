"""Validate the interface layer, write the final workbook, and emit the audit report."""

from __future__ import annotations

import argparse
import os
import shutil
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any, Iterable

from common import (
    AUDIT_REPORT,
    OUTPUT_WORKBOOK,
    PRIMARY_WORKBOOK,
    PROFILE_ID,
    RAW_PORT_COLUMNS,
    RAW_PORT_SHEET,
    REPORT_DIR,
    TARGET_REVIEW_CODES,
    WORK_DIR,
    read_json,
    write_json,
)
from rules import MANUAL_REVIEW_DECISIONS, SEMANTIC_WARNINGS


ERROR_TOKENS = {
    "#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A", "#NUM!", "#NULL!", "#SPILL!", "#CALC!"
}
NEW_SHEETS = [
    "Architecture_Profile",
    "PortType库",
    "InterfaceType库",
    "Compatibility_Rules",
    "PortType映射",
    "Interface_QA",
]


def bool_text(value: bool) -> str:
    return "TRUE" if value else "FALSE"


def table_rows(records: Iterable[dict[str, Any]], columns: list[str]) -> list[list[Any]]:
    return [[record.get(column, "") for column in columns] for record in records]


def apply_table_style(ws, widths: list[float], *, freeze: str = "A2") -> None:
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    max_row = ws.max_row
    max_col = ws.max_column
    if max_row < 1 or max_col < 1:
        return
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(name="Microsoft YaHei", size=10, bold=True, color="FFFFFF")
    body_font = Font(name="Microsoft YaHei", size=10, color="1F1F1F")
    thin = Side(style="thin", color="D9E2F3")
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=thin)
    ws.row_dimensions[1].height = 30
    for row in ws.iter_rows(min_row=2, max_row=max_row, max_col=max_col):
        for cell in row:
            cell.font = body_font
            cell.alignment = Alignment(vertical="top", wrap_text=False)
            cell.border = Border(bottom=Side(style="hair", color="E7E6E6"))
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(1, index).column_letter].width = width
    ws.freeze_panes = freeze
    ws.auto_filter.ref = f"A1:{ws.cell(max_row, max_col).coordinate}"
    ws.sheet_view.showGridLines = False


def highlight_status_rows(ws, status_column: int) -> None:
    from openpyxl.styles import Font, PatternFill

    fills = {
        "FAIL": ("FCE8E6", "B91C1C"),
        "WARNING": ("FFF4CE", "9C6500"),
    }
    for row in range(2, ws.max_row + 1):
        status = ws.cell(row, status_column).value
        if status in fills:
            fill_color, font_color = fills[status]
            for cell in ws[row]:
                cell.fill = PatternFill("solid", fgColor=fill_color)
            ws.cell(row, status_column).font = Font(
                name="Microsoft YaHei", size=10, bold=True, color=font_color
            )


def build_qa_rows(inventory: dict, cleaned: dict, profile: dict, porttypes: dict, interfaces: dict) -> list[dict]:
    leaf_codes = {node["code"] for node in inventory["selection"]["leaves"]}
    item_names = {item["Item中文名"] for item in inventory["item_dictionary"]["items"]}
    item_codes = {item["item_code"] for item in inventory["item_dictionary"]["items"]}
    clean_ports = cleaned["ports"]
    mappings = porttypes["mappings"]
    pt_ids = {row["port_type_id"] for row in porttypes["port_types"]}

    def qa(qa_id: str, passed: bool, description: str, evidence: str) -> dict:
        return {"qa_id": qa_id, "status": "PASS" if passed else "FAIL", "description": description, "evidence": evidence}

    qa_rows = []
    bad_leaf = [p for p in clean_ports if p["code"] not in leaf_codes]
    qa_rows.append(qa("QA-01", not bad_leaf, "所有Clean Raw Port都属于入选叶节点。", f"非法Port数={len(bad_leaf)}；叶节点数={len(leaf_codes)}"))

    bad_items = [p for p in clean_ports if p["交换内容"] not in item_names]
    qa_rows.append(qa("QA-02", not bad_items, "所有交换内容存在Canonical Item。", f"非法Item Port数={len(bad_items)}"))

    mapping_ok = len(mappings) == len(clean_ports)
    if mapping_ok:
        mapping_ok = all(
            mapping["code"] == port["code"]
            and mapping["port_name"] == port["端口名称"]
            and mapping["item_name"] == port["交换内容"]
            for mapping, port in zip(mappings, clean_ports)
        )
    qa_rows.append(qa("QA-03", mapping_ok, "所有Raw Port唯一映射一个PortType。", f"Clean Port={len(clean_ports)}；映射={len(mappings)}"))

    bad_pt_items = [row for row in porttypes["port_types"] if row["item_code"] not in item_codes]
    qa_rows.append(qa("QA-04", not bad_pt_items, "所有PortType引用合法Item。", f"非法引用={len(bad_pt_items)}；PortType={len(pt_ids)}"))

    bad_if_pt = [row for row in interfaces["interface_types"] if row["port_type_id"] not in pt_ids]
    qa_rows.append(qa("QA-05", not bad_if_pt, "所有InterfaceType引用合法PortType。", f"非法引用={len(bad_if_pt)}"))

    bad_signal_direction = [p for p in clean_ports if p["端口类别"] == "信号" and p["方向"] not in {"输入", "输出"}]
    qa_rows.append(qa("QA-06", not bad_signal_direction, "信号方向合法。", f"非法信号方向={len(bad_signal_direction)}"))

    bad_physical_direction = [p for p in clean_ports if p["端口类别"] == "物理" and p["方向"] != "双向物理"]
    qa_rows.append(qa("QA-07", not bad_physical_direction, "物理Port均为双向物理。", f"非法物理方向={len(bad_physical_direction)}"))

    duplicate_count = sum(
        count - 1
        for count in Counter(tuple(port[column] for column in RAW_PORT_COLUMNS) for port in clean_ports).values()
        if count > 1
    )
    qa_rows.append(qa("QA-08", duplicate_count == 0, "没有完全重复Raw Port。", f"重复记录={duplicate_count}"))

    active_mappings = [row for row in mappings if row["profile_active"]]
    active_items = Counter(row["item_name"] for row in active_mappings)
    active_dc_supply = [row for row in active_mappings if row["item_name"] == "直流牵引供电电能"]
    ac_active = active_items["高压交流电能"] > 0
    qa_rows.append(qa("QA-09", ac_active and not active_dc_supply, "当前CRH_AC25KV_SC不存在同时激活的AC/DC互斥供电支路。", f"Active高压交流={active_items['高压交流电能']}；Active直流牵引供电={len(active_dc_supply)}"))

    main_chain = ["高压交流电能", "牵引变压器次级交流电能", "中间直流电能", "三相交流电能", "旋转机械能"]
    missing_main = [item for item in main_chain if active_items[item] < 2]
    qa_rows.append(qa("QA-10", not missing_main, "AC参考实例主能量链具备必要Active Port语义。", "缺失=" + ("无" if not missing_main else "、".join(missing_main))))

    active_pairs = {(row["code"], row["item_name"]) for row in active_mappings}
    regen_required = {
        ("5311", "旋转机械能"),
        ("5311", "三相交流电能"),
        ("5132", "三相交流电能"),
        ("5132", "中间直流电能"),
        ("5131", "中间直流电能"),
        ("5131", "牵引变压器次级交流电能"),
        ("X111", "中间直流电能"),
        ("X111", "储能侧直流电能"),
        ("5142", "中间直流电能"),
        ("5142", "热能/热流"),
    }
    missing_regen = sorted(regen_required - active_pairs)
    qa_rows.append(qa("QA-11", not missing_regen, "再生链具备网侧回馈、超级电容和制动电阻所需接口语义。", f"缺失锚点={missing_regen or '无'}"))

    storage_required = {
        ("X111", "中间直流电能"),
        ("X111", "储能侧直流电能"),
        ("X121", "储能侧直流电能"),
        ("X121", "超级电容SOC"),
    }
    missing_storage = sorted(storage_required - active_pairs)
    qa_rows.append(qa("QA-12", not missing_storage, "储能链具备中间直流、双向DC/DC、储能侧直流与超级电容语义。", f"缺失锚点={missing_storage or '无'}"))

    qa_rows.append(qa("QA-13", True, "不存在非法公式错误和损坏Sheet。", "保存后回读复核待执行"))
    qa_rows.append(qa("QA-14", True, "本轮没有生成任何Connection。", "只生成Raw Port、类型、Profile和规则Sheet"))
    return qa_rows


def scan_formula_errors(workbook_path: Path) -> list[str]:
    from openpyxl import load_workbook

    errors = []
    wb = load_workbook(workbook_path, read_only=False, data_only=False)
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                value = cell.value
                if cell.data_type == "e" or (isinstance(value, str) and any(token in value for token in ERROR_TOKENS)):
                    errors.append(f"{ws.title}!{cell.coordinate}={value}")
    wb.close()
    return errors


def populate_workbook(temp_path: Path, inventory: dict, cleaned: dict, profile: dict, porttypes: dict, interfaces: dict, qa_rows: list[dict]) -> list[str]:
    from openpyxl import load_workbook
    from openpyxl.styles import PatternFill
    from openpyxl.workbook.properties import CalcProperties

    shutil.copy2(PRIMARY_WORKBOOK, temp_path)
    wb = load_workbook(temp_path)
    original_sheets = list(wb.sheetnames)
    if RAW_PORT_SHEET not in wb.sheetnames:
        raise ValueError(f"输出基线缺少{RAW_PORT_SHEET}")
    for sheet_name in NEW_SHEETS:
        if sheet_name in wb.sheetnames:
            del wb[sheet_name]

    raw_ws = wb[RAW_PORT_SHEET]
    if raw_ws.max_row > 1:
        raw_ws.delete_rows(2, raw_ws.max_row - 1)
    for row in table_rows(cleaned["ports"], RAW_PORT_COLUMNS):
        raw_ws.append(row)
    apply_table_style(raw_ws, [12, 24, 24, 20, 36, 12, 12, 30, 20])

    profile_columns = [
        "profile_id", "profile_name", "object_type", "code", "port_name", "item", "active", "reason", "applicability", "object_name"
    ]
    ws = wb.create_sheet("Architecture_Profile")
    ws.append(profile_columns)
    for record in profile["profile_rows"]:
        ws.append([record[column] for column in profile_columns])
    apply_table_style(ws, [20, 34, 14, 12, 38, 26, 12, 60, 18, 26])
    for row in range(2, ws.max_row + 1):
        if ws.cell(row, 7).value is False:
            for cell in ws[row]:
                cell.fill = PatternFill("solid", fgColor="FCE8E6")

    pt_columns = [
        "port_type_id", "port_type_name", "item_code", "item_name", "port_category", "datatype", "unit_medium", "semantic_domain", "semantic_definition"
    ]
    ws = wb.create_sheet("PortType库")
    ws.append(pt_columns)
    for row in table_rows(porttypes["port_types"], pt_columns):
        ws.append(row)
    apply_table_style(ws, [18, 30, 18, 28, 14, 20, 22, 30, 70])

    if_columns = [
        "interface_type_id", "interface_name", "port_type_id", "item_code", "item_name", "interface_domain", "source_category", "compatibility_rule", "datatype", "unit_medium", "description"
    ]
    ws = wb.create_sheet("InterfaceType库")
    ws.append(if_columns)
    for row in table_rows(interfaces["interface_types"], if_columns):
        ws.append(row)
    apply_table_style(ws, [18, 30, 18, 18, 28, 30, 16, 70, 20, 22, 70])

    cr_columns = [
        "rule_id", "applies_to", "constraint_name", "source_requirement", "target_requirement", "predicate", "severity", "description"
    ]
    ws = wb.create_sheet("Compatibility_Rules")
    ws.append(cr_columns)
    for row in table_rows(interfaces["compatibility_rules"], cr_columns):
        ws.append(row)
    apply_table_style(ws, [18, 14, 28, 28, 34, 62, 12, 60])

    mapping_columns = [
        "code", "name", "port_name", "direction", "port_category", "item_code", "item_name", "port_type_id", "interface_type_id", "profile_active"
    ]
    ws = wb.create_sheet("PortType映射")
    ws.append(mapping_columns)
    for row in table_rows(porttypes["mappings"], mapping_columns):
        ws.append(row)
    apply_table_style(ws, [12, 26, 38, 12, 14, 18, 28, 18, 18, 16])
    for row in range(2, ws.max_row + 1):
        if ws.cell(row, 10).value is False:
            for cell in ws[row]:
                cell.fill = PatternFill("solid", fgColor="FFF4CE")

    qa_columns = ["qa_id", "status", "description", "evidence"]
    ws = wb.create_sheet("Interface_QA")
    ws.append(qa_columns)
    for row in table_rows(qa_rows, qa_columns):
        ws.append(row)
    for index, warning in enumerate(SEMANTIC_WARNINGS, start=1):
        ws.append([f"WARNING-{index:02d}", "WARNING", "工程语义待后续实例/拓扑确认。", warning])
    apply_table_style(ws, [18, 14, 62, 100])
    highlight_status_rows(ws, 2)

    if wb.calculation is None:
        wb.calculation = CalcProperties(calcMode="auto", fullCalcOnLoad=True, forceFullCalc=True)
    else:
        wb.calculation.fullCalcOnLoad = True
        wb.calculation.forceFullCalc = True
        wb.calculation.calcMode = "auto"
    wb.save(temp_path)
    wb.close()
    return original_sheets


def finalize_reopen_qa(temp_path: Path, original_sheets: list[str], qa_rows: list[dict]) -> list[str]:
    from openpyxl import load_workbook

    errors = scan_formula_errors(temp_path)
    wb = load_workbook(temp_path, read_only=False, data_only=False)
    missing_original = [name for name in original_sheets if name not in wb.sheetnames]
    forbidden_sheets = [name for name in wb.sheetnames if "Connection" in name or "连接" in name]
    required_sheets = [RAW_PORT_SHEET, *NEW_SHEETS]
    missing_required = [name for name in required_sheets if name not in wb.sheetnames]

    qa13 = next(row for row in qa_rows if row["qa_id"] == "QA-13")
    qa13["status"] = "PASS" if not errors and not missing_original and not missing_required else "FAIL"
    qa13["evidence"] = (
        f"公式错误={len(errors)}；缺失原Sheet={missing_original or '无'}；"
        f"缺失必需Sheet={missing_required or '无'}；工作簿回读成功"
    )
    qa14 = next(row for row in qa_rows if row["qa_id"] == "QA-14")
    qa14["status"] = "PASS" if not forbidden_sheets else "FAIL"
    qa14["evidence"] = f"Connection类Sheet={forbidden_sheets or '无'}"

    ws = wb["Interface_QA"]
    by_id = {ws.cell(row, 1).value: row for row in range(2, ws.max_row + 1)}
    for record in qa_rows:
        row = by_id[record["qa_id"]]
        ws.cell(row, 2, record["status"])
        ws.cell(row, 4, record["evidence"])
    highlight_status_rows(ws, 2)
    wb.save(temp_path)
    wb.close()
    return errors


def write_audit_report(
    inventory: dict,
    cleaned: dict,
    profile: dict,
    porttypes: dict,
    interfaces: dict,
    qa_rows: list[dict],
    report_path: Path,
) -> None:
    ports_by_code: dict[str, list[dict]] = defaultdict(list)
    for port in cleaned["ports"]:
        ports_by_code[port["code"]].append(port)
    used_items = {port["交换内容"] for port in cleaned["ports"]}
    inactive_states = [state for state in profile["port_states"] if not state["active"]]
    report: list[str] = [
        "# Interface Layer 审计报告",
        "",
        f"Profile：`{PROFILE_ID}`",
        "",
        "本轮止于 Clean Raw Port → PortType → InterfaceType → Compatibility Rules；未生成任何 Connection、连接候选、SysML/XMI/SSD，也未修改 Simulink。",
        "",
        "## 汇总",
        "",
        "| 指标 | 数量 |",
        "|---|---:|",
        f"| 输入节点数（入选四级叶节点） | {inventory['selection']['leaf_node_count']} |",
        f"| 清洗前Raw Port数 | {cleaned['before_count']} |",
        f"| 清洗后Raw Port数 | {cleaned['after_count']} |",
        f"| 删除Port数 | {cleaned['deleted_count']} |",
        f"| 修改Port数 | {cleaned['modified_count']} |",
        f"| 新增Port数 | {cleaned['added_count']} |",
        f"| Canonical Item总数 | {inventory['item_dictionary']['item_count']} |",
        f"| 实际使用Item数 | {len(used_items)} |",
        f"| PortType数量 | {len(porttypes['port_types'])} |",
        f"| InterfaceType数量 | {len(interfaces['interface_types'])} |",
        f"| Compatibility Rule数量 | {len(interfaces['compatibility_rules'])} |",
        f"| Active Port数量 | {profile['active_port_count']} |",
        f"| Inactive Variant Port数量 | {profile['inactive_port_count']} |",
        "",
        "## 删除记录",
        "",
    ]
    if cleaned["deleted"]:
        report.extend(["| code | name | port_name | item | 删除原因 |", "|---|---|---|---|---|"])
        for row in cleaned["deleted"]:
            report.append(f"| {row['code']} | {row['name']} | {row['端口名称']} | {row['交换内容']} | {row['reason']} |")
    else:
        report.append("无。")
    report.extend(["", "## 修改记录", ""])
    if cleaned["modified"]:
        for index, row in enumerate(cleaned["modified"], start=1):
            report.append(f"{index}. {row['reason']}：`{row['before']}` → `{row['after']}`")
    else:
        report.append("无字段修改。")
    report.extend(["", "## 新增记录", "", "无。" if not cleaned["added"] else "已按确定性新增规则记录。", "", "## AC/DC相关审查与处理", ""])
    report.extend([
        "- 4235高压电压互感器：判定为AC适用，删除人为添加的1个DC物理Port。",
        "- 4236高压电流互感器：判定为AC适用，删除人为添加的2个DC物理Port。",
        "- 4511主变压器主体：判定为AC适用，原始数据已仅含高压交流与牵引变压器次级交流物理Port，无需修改。",
        "- 4111碳滑板：判定为AC适用并在当前Profile激活。",
        "- 4121受流器：判定为DC适用候选，通用Raw Port保留，但当前Profile整节点停用。",
        "- 4212非真空隔离开关：判定为AC/DC Variant；AC Port激活，2个DC Port停用。",
        "- 5121预充电组件、5122输入组件：判定为AC/DC Variant；各自DC供电候选Port停用，AC链Port激活。",
        "- 5131网侧变流模块：判定为AC适用；5132电机侧变流模块位于中间直流下游，判定为制式无关。",
        f"- 当前Profile中Active直流牵引供电Port为0；Inactive Variant Port共{len(inactive_states)}个。",
        "",
        "### 当前Profile停用Port",
        "",
        "| code | name | port_name | item | reason |",
        "|---|---|---|---|---|",
    ])
    for row in inactive_states:
        report.append(f"| {row['code']} | {row['name']} | {row['port_name']} | {row['item']} | {row['reason']} |")

    report.extend(["", "## 重点节点逐项复核", "", "| code | 清洗后Port数 | 结论 | 复核说明 |", "|---|---:|---|---|"])
    for code in TARGET_REVIEW_CODES:
        status, note = MANUAL_REVIEW_DECISIONS[code]
        report.append(f"| {code} | {len(ports_by_code[code])} | {status} | {note} |")

    report.extend(["", "## QA", "", "| QA | 状态 | 说明 | 证据 |", "|---|---|---|---|"])
    for row in qa_rows:
        report.append(f"| {row['qa_id']} | {row['status']} | {row['description']} | {row['evidence']} |")

    report.extend(["", "## WARNING", ""])
    for index, warning in enumerate(SEMANTIC_WARNINGS, start=1):
        report.append(f"{index}. {warning}")
    report.extend([
        "",
        "## 确定性ID规则",
        "",
        "- `port_type_id = PT-` + Canonical `item_code` 去掉 `ITM-` 前缀，例如 `ITM-PHY-005 → PT-PHY-005`。",
        "- `interface_type_id = IF-` + Canonical `item_code` 去掉 `ITM-` 前缀，例如 `ITM-CMD-006 → IF-CMD-006`。",
        "- direction是Port实例属性，不参与PortType或InterfaceType ID；Input/Output不会导致类型分裂。",
        "- 不使用UUID或随机数。",
        "",
    ])
    report_path.parent.mkdir(parents=True, exist_ok=True)
    report_path.write_text("\n".join(report), encoding="utf-8")


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--inventory", type=Path, default=WORK_DIR / "input_inventory.json")
    parser.add_argument("--cleaned", type=Path, default=WORK_DIR / "cleaned_ports.json")
    parser.add_argument("--profile", type=Path, default=WORK_DIR / "profile_state.json")
    parser.add_argument("--porttypes", type=Path, default=WORK_DIR / "porttypes.json")
    parser.add_argument("--interfaces", type=Path, default=WORK_DIR / "interfacetypes.json")
    parser.add_argument("--output", type=Path, default=OUTPUT_WORKBOOK)
    parser.add_argument("--report", type=Path, default=AUDIT_REPORT)
    parser.add_argument("--summary", type=Path, default=WORK_DIR / "pipeline_summary.json")
    args = parser.parse_args()

    inventory = read_json(args.inventory)
    cleaned = read_json(args.cleaned)
    profile = read_json(args.profile)
    porttypes = read_json(args.porttypes)
    interfaces = read_json(args.interfaces)
    qa_rows = build_qa_rows(inventory, cleaned, profile, porttypes, interfaces)
    initial_failures = [row for row in qa_rows if row["status"] == "FAIL"]
    if initial_failures:
        raise ValueError(f"生成工作簿前QA失败: {initial_failures}")

    args.output.parent.mkdir(parents=True, exist_ok=True)
    temp_path = args.output.with_name(args.output.stem + ".tmp.xlsx")
    if temp_path.exists():
        temp_path.unlink()
    original_sheets = populate_workbook(temp_path, inventory, cleaned, profile, porttypes, interfaces, qa_rows)
    formula_errors = finalize_reopen_qa(temp_path, original_sheets, qa_rows)
    final_failures = [row for row in qa_rows if row["status"] == "FAIL"]
    if final_failures:
        raise ValueError(f"最终QA失败: {final_failures}; 公式错误={formula_errors[:10]}")

    # Final re-open proves the second save remains a valid workbook.
    from openpyxl import load_workbook
    final_check = load_workbook(temp_path, read_only=True, data_only=False)
    final_sheet_names = list(final_check.sheetnames)
    final_check.close()
    if any(name not in final_sheet_names for name in [RAW_PORT_SHEET, *NEW_SHEETS]):
        raise ValueError("最终回读缺少必需Sheet")
    os.replace(temp_path, args.output)
    write_audit_report(inventory, cleaned, profile, porttypes, interfaces, qa_rows, args.report)

    overall = "FAIL" if final_failures else ("WARNING" if SEMANTIC_WARNINGS else "PASS")
    summary = {
        "input_node_count": inventory["selection"]["leaf_node_count"],
        "before_port_count": cleaned["before_count"],
        "after_port_count": cleaned["after_count"],
        "deleted_port_count": cleaned["deleted_count"],
        "modified_port_count": cleaned["modified_count"],
        "added_port_count": cleaned["added_count"],
        "active_port_count": profile["active_port_count"],
        "inactive_variant_port_count": profile["inactive_port_count"],
        "port_type_count": len(porttypes["port_types"]),
        "interface_type_count": len(interfaces["interface_types"]),
        "compatibility_rule_count": len(interfaces["compatibility_rules"]),
        "qa_overall": overall,
        "qa_pass_count": sum(row["status"] == "PASS" for row in qa_rows),
        "qa_fail_count": sum(row["status"] == "FAIL" for row in qa_rows),
        "warnings": SEMANTIC_WARNINGS,
        "output_workbook": str(args.output.resolve()),
        "code_directory": str(Path(__file__).resolve().parent),
        "audit_report": str(args.report.resolve()),
    }
    write_json(args.summary, summary)
    print(f"QA总体: {overall}")
    print(f"输出Excel: {args.output.resolve()}")
    print(f"审计报告: {args.report.resolve()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
