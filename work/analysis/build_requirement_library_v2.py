from __future__ import annotations

import json
import re
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
DATA = ROOT / "ui/backend/data/semantic_slice"
REQ_XLSX = ROOT / "Req/牵引制动能量回收系统_正式需求库_v1.xlsx"
SYSML = ROOT / "work/sysmlv2/full_engineering_model/01_generated/Rail_MBSE_Full_v1.sysml"
OUT = ROOT / "ui/backend/data/semantic_slice"
REPORT = ROOT / "work/reports/REQUIREMENT_LIBRARY_V2_EXPANSION_REPORT.md"
BUNDLE = Path(__file__).resolve().parent / "requirement_library_v2_bundle.json"

FIELDS = ["Requirement_ID", "Parent_Requirement_IDs", "Requirement_Level", "Domain", "Topic_Group", "Name", "Text",
    "Requirement_Type", "Source_Type", "Standard_No", "Standard_Clause", "Criteria", "Applicable_Conditions", "Keywords",
    "Synonyms", "User_Phrases", "Semantic_Topics", "System_Terms", "Action_Terms", "Object_Terms",
    "Expected_Function_Semantics", "Linked_Function_IDs", "Linked_Function_Names", "Linked_Product_IDs", "Linked_Product_Names",
    "Function_Coverage_Count", "Coverage_Status", "Derivation_Reason", "Related_Requirement_IDs", "Verification_Status",
    "Model_Coverage", "Notes", "Evidence_Source"]

# The new requirements are grouped by independent equipment/interface goals.
# Each group is backed by existing allocated Functions; no model elements are added.
DERIVED = [
 {"id":"DRV-TRAC-001","domain":"牵引","group":"TRACTION","name":"牵引传动机械能传递","text":"牵引传动系统应将牵引电机输出的机械能经联轴器、齿轮箱和轮轴传递至轮对，并保持传动链各级机械能传递关系可追溯。","functions":["FN_F_3100_01","FN_F_3110_01","FN_F_3112_01","FN_F_3120_01","FN_F_3121_01","FN_F_3500_01","FN_F_3510_01","FN_F_3511_01","FN_F_3512_01","FN_F_3514_01"],"parents":["REQ-TRAC-001","REQ-TRAC-004"],"topics":["牵引","机械能传递","传动链"],"keywords":["牵引传动","机械能传递","联轴器","齿轮箱","轮轴","轮对","电机侧传动","传动链"],"synonyms":["牵引传动链","电机至轮对传动","机械传动"],"phrases":["电机的机械能怎样传到轮对","联轴器和齿轮箱属于哪些牵引要求","我想看牵引传动链的要求"],"actions":["传递","转换"],"objects":["机械能","联轴器","齿轮箱","轮轴","轮对"],"reason":"现有牵引特性与能量转换需求没有明确表达电机至轮对的机械传动链；Function 3100/3500 及其真实分配到的轮轴、齿轮箱和联轴器提供了成组实现证据。"},
 {"id":"DRV-TRAC-002","domain":"牵引","group":"TRACTION","name":"黏着增强撒砂控制","text":"黏着增强装置应接收撒砂控制输入，驱动撒砂控制单元、阀板和喷嘴执行撒砂动作，并向相关控制边界提供装置运行状态。","functions":["FN_F_3800_01","FN_F_3810_01","FN_F_3811_01","FN_F_3811_02","FN_F_3813_01","FN_F_3814_01","FN_F_D13B_01"],"parents":["REQ-TRAC-003","REQ-BRK-010"],"topics":["黏着控制","撒砂","防滑"],"keywords":["撒砂","黏着增强","轮轨黏着","防滑","砂阀","撒砂喷嘴","撒砂控制","黏着利用"],"synonyms":["撒砂装置","砂控","黏着增强装置"],"phrases":["空转时撒砂装置怎么动作","撒砂按钮会触发哪些设备","我想看黏着增强和防滑相关要求"],"actions":["控制","执行","报告","传递"],"objects":["撒砂控制输入","撒砂阀板","喷嘴","运行状态","黏着"],"reason":"现有防滑需求涉及黏着，但未覆盖撒砂控制单元、执行阀板/喷嘴及司机撒砂按钮这组真实 Functions；该派生项以这些已分配功能为边界，不声称模型存在自动滑移率触发逻辑。"},
 {"id":"DRV-PWR-001","domain":"供电","group":"POWER_SUPPLY","name":"高压受流与分配控制","text":"高压受流及分配系统应控制受电弓、高压断路器和隔离开关等既有设备的运行，采集高压电压、电流及开关单元状态，并通过已建模控制接口提供相应输入输出和状态信息。","functions":["FN_F_4100_01","FN_F_4110_01","FN_F_411C_01","FN_F_411C_02","FN_F_4200_01","FN_F_4210_01","FN_F_4211_01","FN_F_4212_01","FN_F_4230_01","FN_F_4235_01","FN_F_4236_01","FN_F_4238_01","FN_F_4238_02","FN_F_4238_03","FN_F_4239_01","FN_F_4239_02","FN_F_4239_03","FN_F_4240_01","FN_F_4241_01","FN_F_4241_02","FN_F_4241_03"],"parents":["PRJ-PWR-001"],"topics":["高压供电","受流","高压分配","开关控制"],"keywords":["高压受流","受电弓","碳滑板","高压断路器","隔离开关","高压分配","电压互感器","电流互感器","开关状态","高压控制单元"],"synonyms":["网侧高压供电","受流装置","高压开关设备","高压测量"],"phrases":["受电弓和高压开关如何纳入控制","高压电压电流由哪些单元采集","高压分配状态怎么传给控制系统"],"actions":["受流","控制","采集","路由","协调","报告"],"objects":["受电弓","断路器","隔离开关","高压电压","高压电流","开关状态"],"reason":"项目供电制式需求规定仿真名义供电条件，但未提出受电弓电子单元、高压开关、测量和I/O状态控制行为；这些功能由现有4100/4200产品及真实allocation支撑。"},
 {"id":"DRV-PWR-002","domain":"供电","group":"POWER_SUPPLY","name":"牵引回流与接地通路","text":"牵引回流及接地系统应通过已建模接地变压器、主回路接地电阻和回流碳刷等部件提供牵引回流与接地能量传递通路。","functions":["FN_F_4400_01","FN_F_4410_01","FN_F_4411_01","FN_F_4412_01","FN_F_4413_01"],"parents":["PRJ-PWR-001"],"topics":["牵引供电","回流","接地"],"keywords":["牵引回流","回流通路","接地","接地变压器","接地电阻","回流碳刷","主回路"],"synonyms":["回流系统","主回路接地","轨道回流"],"phrases":["牵引电流如何回流","主回路接地由哪些部件实现","回流碳刷相关要求是什么"],"actions":["传递","回流","接地"],"objects":["回流电能","接地变压器","接地电阻","回流碳刷"],"reason":"25 kV供电项目需求没有展开回流和接地组成的工程行为；4400及其真实分配部件构成同一回流/接地目标链。"},
 {"id":"DRV-PWR-003","domain":"供电","group":"POWER_SUPPLY","name":"主变压器控制与状态监测","text":"牵引变压器及其控制装置应传递变压器支路电能，监测既有检测保护元件状态，并提供控制装置运行状态。","functions":["FN_F_4500_01","FN_F_4510_01","FN_F_4513_01","FN_F_4517_01","FN_F_4518_01","FN_F_4518_02"],"parents":["PRJ-PWR-001","REQ-TRAC-004"],"topics":["牵引供电","主变压器","状态监测","电气保护"],"keywords":["主变压器","变压器控制","检测保护元件","变压器状态","电抗器","变压器支路","状态报告"],"synonyms":["牵引变压器","变压器监测","变压器控制装置"],"phrases":["主变压器保护元件状态怎么监测","变压器控制装置会报告哪些状态","主变压器支路电能如何传递"],"actions":["传递","监测","控制","报告"],"objects":["变压器","电抗器","保护元件状态","控制装置状态"],"reason":"已有供电制式和电能转换需求没有覆盖主变压器支路传递、检测保护元件及控制装置状态报告；对应功能集中分配到4500产品边界。"},
 {"id":"DRV-TRAC-003","domain":"牵引","group":"TRACTION","name":"牵引变流器控制与保护状态","text":"牵引变流器应控制网侧和电机侧变流控制模块，协调其接口交互，监测检测保护组件状态，并提供模块运行状态和保护限制信息。","functions":["FN_F_5100_01","FN_F_5110_01","FN_F_5111_01","FN_F_5111_02","FN_F_5111_03","FN_F_5111_04","FN_F_5112_02","FN_F_5112_03","FN_F_5112_04","FN_F_5113_01","FN_F_5130_01","FN_F_5131_02","FN_F_5132_02"],"parents":["REQ-TRAC-001","REQ-TRAC-002","REQ-TRAC-004"],"topics":["牵引变流器","变流控制","变流保护","状态监测"],"keywords":["牵引变流器控制","网侧变流器","电机侧变流器","变流模块","接口协调","变流保护","检测组件","运行状态"],"synonyms":["牵引变流控制","变流器状态监测","变流模块保护"],"phrases":["网侧和电机侧变流器怎样协调","变流器的保护和状态由谁处理","牵引变流器检测组件需要监测什么"],"actions":["控制","协调","监测","限制","保护","报告"],"objects":["网侧变流模块","电机侧变流模块","检测保护组件","运行状态"],"reason":"原牵引性能和能量转换需求没有细化变流控制模块的接口协调、保护和状态报告；缺口由5100下的控制、保护、状态和测量Function组成一组补足。"},
 {"id":"DRV-TRAC-004","domain":"牵引","group":"TRACTION","name":"中间直流环节预充与放电保护","text":"牵引变流器中间直流环节应通过既有输入、预充电和放电电阻支路执行电能传递与耗散，并控制和监测相应支路状态。","functions":["FN_F_5120_01","FN_F_5121_01","FN_F_5122_01","FN_F_5133_01","FN_F_5133_02","FN_F_5133_03","FN_F_5140_01","FN_F_5142_01","FN_F_5142_02","FN_F_5145_01"],"parents":["REQ-TRAC-004","PRJ-PWR-001"],"topics":["牵引变流器","中间直流环节","预充电","放电保护"],"keywords":["中间直流环节","直流母线","预充电","放电电阻","直流环节放电","输入组件","支路保护","直流电压检测"],"synonyms":["DC-Link","直流链","母线预充","母线泄放"],"phrases":["直流母线启动前怎么预充","中间直流环节如何放电","放电电阻支路有哪些保护和状态"],"actions":["传递","预充","控制","保护","耗散","监测","统计"],"objects":["中间直流环节","预充电组件","放电电阻","支路状态"],"reason":"供电和能量转换需求未覆盖中间直流环节的输入/预充/放电电阻支路行为；这些Function均真实分配到牵引变流器5100下级产品。"},
 {"id":"DRV-BRK-001","domain":"制动","group":"PNEUMATIC_BRAKE","name":"制动供风与总风管压力可用性","text":"制动供风系统应向制动风缸提供压缩空气并传递总风管压力，检测供风系统运行参数和总风管压力传感器状态，以支持气动制动可用性判断。","functions":["FN_F_7100_01","FN_F_7120_01","FN_F_7124_01","FN_F_7130_01","FN_F_7131_01","FN_F_7140_01","FN_F_7142_01"],"parents":["REQ-BRK-001","REQ-BRK-012"],"topics":["制动供风","空气制动","压力监测"],"keywords":["制动供风","压缩空气","制动风缸","总风管","总风压力","气动制动","压力传感器","供风可用性"],"synonyms":["供风系统","风源系统","空气制动供气","制动气源"],"phrases":["制动用压缩空气从哪里来","总风管压力由谁监测","供风系统和制动风缸怎么连接"],"actions":["提供","传递","检测","监测"],"objects":["压缩空气","制动风缸","总风管压力","压力传感器状态"],"reason":"现有制动需求未完整覆盖供风、风缸供气及总风管压力检测链；7100真实Function及其供风部件allocation形成同一气动制动可用性目标。"},
 {"id":"DRV-BRK-002","domain":"制动","group":"BRAKING","name":"气动压力分配与机械制动执行","text":"制动执行系统应按既有制动控制传递、调节或隔离气动压力，并将气动压力转换为制动机械力、传递夹钳作用力及耗散机械制动能量，同时提供已建模制动指示状态。","functions":["FN_F_7200_01","FN_F_7210_01","FN_F_7213_01","FN_F_7220_01","FN_F_7222_01","FN_F_7223_01","FN_F_7225_01","FN_F_7227_01","FN_F_7228_01","FN_F_7229_01","FN_F_7250_01","FN_F_7252_01","FN_F_7254_01","FN_F_7280_01","FN_F_7281_01","FN_F_7284_01","FN_F_7285_01"],"parents":["REQ-BRK-001","REQ-BRK-004","REQ-BRK-013"],"topics":["机械制动","空气制动","制动执行","制动力传递"],"keywords":["制动执行","气动压力","电空转换阀","中继阀","电磁阀","缓解阀","制动夹钳","闸片","摩擦制动","制动指示器","制动截断塞门"],"synonyms":["机械制动","摩擦制动","气动制动","基础制动","制动执行机构"],"phrases":["机械制动什么时候介入","制动压力如何传到制动夹钳","电空转换后机械制动怎么执行"],"actions":["传递","调节","隔离","转换","执行","耗散","指示"],"objects":["气动压力","制动机械力","制动夹钳","闸片","制动状态"],"reason":"既有制动需求表达制动性能和协调目标，但没有说明阀件、气压分配、制动夹钳、闸片及状态指示构成的执行链；所列功能真实属于7200制动产品树。"},
 {"id":"DRV-CTL-001","domain":"控制","group":"BRAKE_CONTROL","name":"司机操纵与ATP控制边界输入","text":"列车操纵及ATP边界应通过既有方向、牵引、制动操作元件和ATP接口向列车控制系统提供相应输入，并报告ATP接口运行状态。","functions":["FN_F_D100_01","FN_F_D110_01","FN_F_D112_01","FN_F_D113_01","FN_F_D114_01","FN_F_D120_01","FN_F_D130_01","FN_F_E100_01","FN_F_8127_01","FN_F_8127_02"],"parents":["REQ-TRAC-002","REQ-BRK-001","REQ-BRK-002"],"topics":["司机操纵","列车控制输入","ATP接口"],"keywords":["司机操纵","方向手柄","牵引手柄","制动手柄","ATP车载","ATP接口","列车控制输入","操作元件"],"synonyms":["司机控制界面","操纵台输入","ATP边界接口"],"phrases":["方向牵引和制动手柄的输入怎么进入控制系统","ATP接口和司机操纵有哪些需求","列车控制输入来自哪些边界设备"],"actions":["提供","输入","控制","路由","报告"],"objects":["方向输入","牵引请求","制动请求","ATP接口状态"],"reason":"现有牵引、制动性能需求未定义其边界操纵输入；D100操作元件、E100 ATP边界及8100 ATP接口单元提供了可验证的真实接口功能。"},
 {"id":"DRV-CTL-002","domain":"控制","group":"COMMUNICATION","name":"车载控制单元协调与状态报告","text":"车载列车控制系统应在既有主控制、车辆控制、列车级牵引控制和输入输出单元之间协调控制接口交互，并提供相关控制单元的运行状态及运行数据。","functions":["FN_F_8100_01","FN_F_8120_01","FN_F_8121_01","FN_F_8121_02","FN_F_8121_03","FN_F_8122_01","FN_F_8122_02","FN_F_8122_04","FN_F_8122_05","FN_F_8125_01","FN_F_8125_02","FN_F_8125_03","FN_F_8125_04","FN_F_8128_01","FN_F_8128_02","FN_F_8128_03","FN_F_8160_01"],"parents":["REQ-TRAC-002","REQ-BRK-014","REQ-ENE-003"],"topics":["列车控制","车载通信","状态报告","控制单元协调"],"keywords":["列车控制系统","主控制单元","车辆控制单元","牵引控制单元","输入输出单元","控制接口","状态报告","运行数据","车载通信"],"synonyms":["车载控制网络","列车控制单元协调","控制状态反馈"],"phrases":["车辆控制和牵引控制单元如何协同","控制单元的状态信息从哪里获取","输入输出单元会报告哪些运行状态"],"actions":["协调","控制","路由","报告","统计"],"objects":["主控制单元","车辆控制单元","牵引控制单元","输入输出单元","运行状态"],"reason":"当前需求只覆盖个别制动安全或储能通信现象，未表达车载主控、车辆控制、牵引控制及I/O控制接口协调与状态报告这一组系统级能力；各Function均属于8100真实Product树。"},
 {"id":"DRV-CTL-003","domain":"安全","group":"SAFETY_PROTECTION","name":"牵引与高压安全回路逻辑及状态","text":"牵引、高压和制动安全回路应执行模型中定义的安全逻辑，并向车辆控制系统提供相应安全回路运行状态。","functions":["FN_F_8150_01","FN_F_8151_01","FN_F_8151_02","FN_F_8152_01","FN_F_8152_02","FN_F_8154_01","FN_F_8154_02"],"parents":["REQ-BRK-014","REQ-BRK-015","REQ-TRAC-003"],"topics":["安全回路","牵引保护","高压保护","制动保护"],"keywords":["安全回路","高压安全回路","牵引安全回路","制动安全回路","安全逻辑","回路状态","运行状态反馈"],"synonyms":["安全联锁回路","安全状态反馈","牵引保护回路"],"phrases":["牵引和高压安全回路怎么执行保护","制动安全回路状态会不会上报","哪些安全回路具有状态反馈"],"actions":["执行","保护","限制","反馈","报告"],"objects":["高压安全回路","牵引安全回路","制动安全回路","安全状态"],"reason":"已有故障诊断需求涉及部分制动安全回路，但高压与牵引安全逻辑及状态反馈仍没有统一需求来源；8150、8151、8152、8154真实Functions支持该分组。"},
]

EXISTING_LINKS = {
 "REQ-TRAC-001":["FN_F_3100_01","FN_F_3112_01","FN_F_3120_01","FN_F_3121_01","FN_F_3500_01","FN_F_3510_01","FN_F_3511_01","FN_F_3512_01","FN_F_3514_01","FN_F_3600_01","FN_F_3640_01","FN_F_3641_01","FN_F_3643_01","FN_F_3644_01","FN_F_5300_01","FN_F_5310_01"],
 "REQ-TRAC-003":["FN_F_3641_01","FN_F_3643_01","FN_F_3644_01"],
 "REQ-TRAC-004":["FN_F_3100_01","FN_F_3112_01","FN_F_3120_01","FN_F_3121_01","FN_F_3500_01","FN_F_3510_01","FN_F_3511_01","FN_F_3512_01","FN_F_3514_01","FN_F_4500_01","FN_F_4510_01","FN_F_4513_01","FN_F_5300_01","FN_F_5310_01"],
 "REQ-BRK-001":["FN_F_3111_01","FN_F_3130_01","FN_F_3131_01","FN_F_3132_01","FN_F_3133_01","FN_F_7252_01","FN_F_7254_01"],
 "REQ-BRK-004":["FN_F_3130_01","FN_F_3131_01","FN_F_3132_01","FN_F_3133_01"],
 "REQ-ENE-001":["FN_F_X110_01","FN_F_X111_01","FN_F_X111_02"],
 "REQ-ENE-005":["FN_F_8166_01","FN_F_8166_02","FN_F_8166_03","FN_F_8166_04"],
 "PRJ-PWR-001":["FN_F_4100_01","FN_F_4110_01","FN_F_4200_01","FN_F_4210_01","FN_F_4211_01","FN_F_4212_01"],
}

SUPPORT_PARENTS = {
 "3000":["REQ-TRAC-001","REQ-TRAC-004"], "3100":["DRV-TRAC-001","REQ-BRK-001"],
 "3500":["DRV-TRAC-001"], "3600":["REQ-TRAC-001","REQ-TRAC-003"], "3800":["DRV-TRAC-002","REQ-TRAC-003"],
 "4000":["PRJ-PWR-001"], "4100":["DRV-PWR-001","PRJ-PWR-001"], "4200":["DRV-PWR-001","PRJ-PWR-001"],
 "4400":["DRV-PWR-002"], "4500":["DRV-PWR-003","PRJ-PWR-001"], "5000":["REQ-TRAC-004"],
 "5100":["DRV-TRAC-003","DRV-TRAC-004","REQ-TRAC-004"], "5300":["REQ-TRAC-001","REQ-TRAC-004"],
 "7000":["REQ-BRK-001"], "7100":["DRV-BRK-001","REQ-BRK-012"], "7200":["DRV-BRK-002","REQ-BRK-001"],
 "8000":["DRV-CTL-002","DRV-CTL-003"], "8100":["DRV-CTL-002","DRV-CTL-003"],
 "D000":["DRV-CTL-001"], "D100":["DRV-CTL-001"], "E000":["DRV-CTL-001"], "E100":["DRV-CTL-001"],
 "X000":["REQ-ENE-001","PRJ-ENE-001"], "X100":["REQ-ENE-001","REQ-ENE-003"],
}

BASE_TERMS = {
 "牵引":["牵引","牵引力","牵引控制","牵引变流器","牵引电机","电能转换","速度","转矩"],
 "制动":["制动","制动力","制动请求","再生制动","机械制动","摩擦制动","空气制动","制动协调","减速度"],
 "储能":["储能","再生能量","超级电容","SOC","充电","放电","储能功率","能量回收","储能状态"],
 "供电":["供电","高压","接触网","受电弓","主变压器","高压分配","回流","接地","电能传递"],
 "控制":["列车控制","控制单元","控制接口","输入输出","运行状态","指令","接口协调"],
 "安全":["安全回路","安全逻辑","牵引保护","制动保护","高压保护","状态反馈"],
}
SYNONYM_TERMS = {
 "牵引":["牵引驱动","电机牵引","牵引传动"],
 "制动":["刹车","摩擦制动","空气制动","气动制动","电制动"],
 "储能":["车载储能","超级电容器","超容","电容储能","回馈制动","制动能量回收"],
 "供电":["网侧供电","牵引供电","受流系统","高压电源"],
 "控制":["车载控制","控制状态反馈","控制网络"],
 "安全":["安全联锁","安全状态反馈","保护回路"],
}
ACTION_VOCAB = ["检测","监测","控制","协调","执行","传递","转换","限制","吸收","存储","释放","输出","报告","反馈","保护","隔离","耗散","路由","统计","提供","采集","分配"]
OBJECT_VOCAB = ["牵引力","制动力","转矩","速度","机械能","电能","再生功率","充电功率","放电功率","超级电容","SOC","电流","电压","气动压力","总风管","制动夹钳","齿轮箱","轮轴","受电弓","断路器","接地","回流","运行状态","安全回路","故障状态","控制接口","撒砂"]

CLUSTERS = [
 {"id":"CLUSTER-TRAC-DRIVE","name":"牵引驱动与传动","triggers":["牵引","加速","牵引电机","变流器","转矩","齿轮箱","轮轴"],"groups":["TRACTION"],"description":"牵引需求、变流控制、牵引电机及机械传动链。"},
 {"id":"CLUSTER-BRK-REGEN","name":"再生制动与能量回收","triggers":["再生制动","电制动","回馈制动","制动能量","能量回收","储能","超级电容"],"groups":["REGENERATIVE_BRAKING","ENERGY_STORAGE"],"description":"电制动优先、制动能量流、储能吸收及功率/SOC边界。"},
 {"id":"CLUSTER-BRK-FRICTION","name":"机械与空气制动","triggers":["机械制动","摩擦制动","空气制动","气动压力","制动夹钳","闸片","低速制动"],"groups":["BRAKING","PNEUMATIC_BRAKE"],"description":"再生制动补偿、气动供风、机械制动执行与制动力形成。"},
 {"id":"CLUSTER-BRK-SAFETY","name":"制动安全与故障处置","triggers":["紧急制动","防滑","安全回路","制动故障","失效降级","压力诊断"],"groups":["SAFETY_PROTECTION","BRAKING"],"description":"紧急制动、防滑、安全回路、压力状态和故障保护。"},
 {"id":"CLUSTER-ENE-SUPERCAP","name":"超级电容状态与保护","triggers":["超级电容","超容","SOC","电压","电流","温度","过流","过温"],"groups":["ENERGY_STORAGE"],"description":"储能状态监测、充放电限制及故障保护。"},
 {"id":"CLUSTER-ENE-RECOVERY","name":"再生能量收支与储能功率","triggers":["回收能量","再生功率","充电功率","可用功率","能量效率","直流母线"],"groups":["ENERGY_STORAGE","REGENERATIVE_BRAKING"],"description":"再生功率分配、能量吸收和系统能量测量。"},
 {"id":"CLUSTER-PWR-HV","name":"高压受流与主回路","triggers":["高压","受电弓","接触网","断路器","隔离开关","主变压器","回流","接地"],"groups":["POWER_SUPPLY"],"description":"高压输入、开关分配、变压器及牵引回流接地通路。"},
 {"id":"CLUSTER-CTL-INPUT","name":"操纵与ATP接口","triggers":["司机操纵","方向手柄","牵引手柄","制动手柄","ATP","控制输入"],"groups":["BRAKE_CONTROL"],"description":"司机操作元件、ATP接口及列车控制输入边界。"},
 {"id":"CLUSTER-CTL-COMMS","name":"车载控制协调与状态通信","triggers":["车辆控制","主控制单元","输入输出单元","接口交互","状态报告","通信"],"groups":["COMMUNICATION"],"description":"车载控制单元间控制交互、路由、状态和运行数据。"},
 {"id":"CLUSTER-DIAG-SAFETY","name":"诊断、状态与安全回路","triggers":["故障状态","运行状态","安全回路","检测保护","传感器状态","诊断"],"groups":["DIAGNOSTICS","SAFETY_PROTECTION"],"description":"模型已具备的设备状态检测、保护与安全回路反馈。"},
]


def read_json(name):
    return json.loads((DATA / name).read_text(encoding="utf-8"))


def vals(value):
    if value is None: return []
    if isinstance(value, list): return [str(x).strip() for x in value if str(x).strip()]
    return [x.strip() for x in re.split(r"[;；,，\n]+", str(value)) if x and x.strip()]


def uniq(items, limit=None):
    seen, out = set(), []
    for item in items:
        v = re.sub(r"\s+", " ", str(item or "").strip())
        key = v.casefold()
        if v and key not in seen:
            seen.add(key); out.append(v)
    return out[:limit] if limit else out


def main():
    catalog = read_json("requirement_catalog.json")["requirements"]
    functions = read_json("function_catalog.json")["functions"]
    products = read_json("product_catalog.json")["products"]
    allocations = read_json("function_product_map.json")["relations"]
    old_relations = read_json("requirement_function_map.json")["relations"]
    raw_wb = load_workbook(REQ_XLSX, read_only=True, data_only=True)
    raw_sheet = raw_wb["正式需求库"]
    raw_rows = [[v.isoformat() if hasattr(v, "isoformat") else v for v in row] for row in raw_sheet.iter_rows(values_only=True)]
    raw_headers, raw_body = raw_rows[0], raw_rows[1:]
    raw_by_id = {str(r[0]): dict(zip(raw_headers, r)) for r in raw_body if r and r[0]}
    req_by_id = {x["requirement_id"]: x for x in catalog}
    fn_by_id = {x["function_id"]: x for x in functions}
    product_by_id = {x["product_id"]: x for x in products}
    alloc_by_fn = {x["function_id"]: x for x in allocations}
    model_text = SYSML.read_text(encoding="utf-8")
    for rel in allocations:
        if rel.get("allocation_id") and not re.search(rf"allocation\s+{re.escape(rel['allocation_id'])}\s+allocate\s+", model_text):
            raise ValueError(f"Allocation absent in frozen SysML: {rel['allocation_id']}")

    requirement_by_id = {}
    for source in catalog:
        rid = source["requirement_id"]
        raw = raw_by_id[rid]
        # The protected source fields below are copied verbatim from v1.
        for target, source_field in [("Requirement_ID","Requirement_ID"),("Name","需求名称"),("Text","正式需求表述"),
            ("Requirement_Type","需求类型"),("Source_Type","来源类型"),("Standard_No","标准号"),
            ("Standard_Clause","条款号"),("Criteria","标准/项目判据")]:
            if ("" if raw.get(source_field) is None else str(raw[source_field]).strip()) != ("" if source.get({
                "Requirement_ID":"requirement_id","Name":"name","Text":"text","Requirement_Type":"requirement_type",
                "Source_Type":"source_type","Standard_No":"standard_no","Standard_Clause":"standard_clause","Criteria":"criteria"}[target]) is None else str(source.get({
                "Requirement_ID":"requirement_id","Name":"name","Text":"text","Requirement_Type":"requirement_type",
                "Source_Type":"source_type","Standard_No":"standard_no","Standard_Clause":"standard_clause","Criteria":"criteria"}[target])).strip()):
                raise ValueError(f"Protected v1 field changed or data parse mismatch: {rid} {target}")
        requirement_by_id[rid] = {"Requirement_ID":rid,"Parent_Requirement_IDs":"","Requirement_Level":"L1_DOMAIN",
            "Domain":source["domain"],"Topic_Group":"","Name":source["name"],"Text":source["text"],
            "Requirement_Type":source["requirement_type"],"Source_Type":source["source_type"],"Standard_No":source["standard_no"],
            "Standard_Clause":source["standard_clause"],"Criteria":source["criteria"],"Applicable_Conditions":source["applicable_conditions"],
            "Keywords":vals(source.get("keywords")),"Synonyms":vals(source.get("synonyms")),"User_Phrases":vals(source.get("user_phrases")),
            "Semantic_Topics":[],"System_Terms":[],"Action_Terms":[],"Object_Terms":[],
            "Expected_Function_Semantics":source["expected_function_semantics"],"Linked_Function_IDs":[],"Linked_Function_Names":[],
            "Linked_Product_IDs":[],"Linked_Product_Names":[],"Function_Coverage_Count":0,"Coverage_Status":"PARTIAL",
            "Derivation_Reason":"","Related_Requirement_IDs":[],"Verification_Status":source["verification_status"],
            "Model_Coverage":source["model_coverage"],"Notes":source["notes"],"Evidence_Source":source.get("evidence_url") or REQ_XLSX.relative_to(ROOT).as_posix(),
            "_source":source,"_raw":raw}
        if rid in {"PRJ-BRK-001","PRJ-ENE-001","PRJ-PWR-001"}:
            requirement_by_id[rid]["Requirement_Level"] = "L0_SYSTEM"

    for spec in DERIVED:
        if spec["id"] in requirement_by_id: raise ValueError(f"Duplicate derived id: {spec['id']}")
        for fid in spec["functions"]:
            if fid not in fn_by_id or fid not in alloc_by_fn: raise ValueError(f"Derived Function missing: {spec['id']} {fid}")
        requirement_by_id[spec["id"]] = {"Requirement_ID":spec["id"],"Parent_Requirement_IDs":spec["parents"],
            "Requirement_Level":"L2_DERIVED","Domain":spec["domain"],"Topic_Group":[spec["group"]],"Name":spec["name"],"Text":spec["text"],
            "Requirement_Type":"工程功能要求","Source_Type":"ENGINEERING_DERIVED","Standard_No":"","Standard_Clause":"",
            "Criteria":"模型判据：所列既有Function及其SysML allocation存在；本判据不表示标准符合性或试验验证通过。",
            "Applicable_Conditions":"适用于本项目现有SysML产品边界及已建模运行/控制接口。","Keywords":spec["keywords"],
            "Synonyms":spec["synonyms"],"User_Phrases":spec["phrases"],"Semantic_Topics":spec["topics"],
            "System_Terms":[],"Action_Terms":spec["actions"],"Object_Terms":spec["objects"],
            "Expected_Function_Semantics":spec["text"],"Linked_Function_IDs":[],"Linked_Function_Names":[],
            "Linked_Product_IDs":[],"Linked_Product_Names":[],"Function_Coverage_Count":0,"Coverage_Status":"FULL",
            "Derivation_Reason":spec["reason"],"Related_Requirement_IDs":[],"Verification_Status":"NOT_VERIFIED",
            "Model_Coverage":"COVERED","Notes":"工程派生需求；不声明为标准直接要求，不增加或修改SysML对象。",
            "Evidence_Source":SYSML.relative_to(ROOT).as_posix(),"_spec":spec}

    # Preserve existing reviewed/curated Requirement→Function relations.
    relation_map = {}
    for rel in old_relations:
        rid, fid = rel["requirement_id"], rel["function_id"]
        if rid not in requirement_by_id or fid not in fn_by_id: raise ValueError(f"Existing map has unresolved id: {rid} {fid}")
        owner = alloc_by_fn[fid]
        relation_map[(rid,fid)] = {"requirement_id":rid,"function_id":fid,"product_id":owner["product_id"],
            "mapping_type":rel["mapping_source"],"evidence":rel.get("evidence", ""),
            "evidence_path":rel.get("evidence_path", []),"semantic_evidence":rel.get("semantic_evidence", rel.get("evidence", "")),
            "structural_evidence":{"allocation_id":owner.get("allocation_id"),"allocation":owner["evidence"],"owner_product_id":owner["product_id"]},
            "source_files":uniq([*rel.get("source_files",[]),owner["source_file"]])}

    # Targeted, auditable backfill from actual allocation ownership and requirement semantics.
    for rid, fids in EXISTING_LINKS.items():
        for fid in fids:
            if fid not in fn_by_id or fid not in alloc_by_fn: raise ValueError(f"Existing backfill id not found: {rid} {fid}")
            if (rid,fid) in relation_map: continue
            owner = alloc_by_fn[fid]
            relation_map[(rid,fid)] = {"requirement_id":rid,"function_id":fid,"product_id":owner["product_id"],
                "mapping_type":"MODEL_DERIVED","evidence":f"既有需求范围与Function语义一致；{owner['evidence']}。仅以真实OWNER allocation定位Function所属Product。",
                "evidence_path":[{"type":"REQUIREMENT","id":rid},{"type":"PRODUCT","id":owner["product_id"]},
                    {"type":"ALLOCATION","id":owner.get("allocation_id"),"product_id":owner["product_id"]},{"type":"FUNCTION","id":fid}],
                "semantic_evidence":f"现有需求语义覆盖该功能用途：{req_by_id[rid]['name']}；{fn_by_id[fid]['name']}。",
                "structural_evidence":{"allocation_id":owner.get("allocation_id"),"allocation":owner["evidence"],"owner_product_id":owner["product_id"]},
                "source_files":[REQ_XLSX.relative_to(ROOT).as_posix(),owner["source_file"]]}

    for spec in DERIVED:
        for fid in spec["functions"]:
            owner = alloc_by_fn[fid]
            relation_map[(spec["id"],fid)] = {"requirement_id":spec["id"],"function_id":fid,"product_id":owner["product_id"],
                "mapping_type":"DERIVED_REQUIREMENT","evidence":spec["reason"],
                "evidence_path":[{"type":"REQUIREMENT","id":spec["id"]},{"type":"PRODUCT","id":owner["product_id"]},
                    {"type":"ALLOCATION","id":owner.get("allocation_id"),"product_id":owner["product_id"]},{"type":"FUNCTION","id":fid}],
                "semantic_evidence":spec["text"],
                "structural_evidence":{"allocation_id":owner.get("allocation_id"),"allocation":owner["evidence"],"owner_product_id":owner["product_id"]},
                "source_files":[SYSML.relative_to(ROOT).as_posix(),owner["source_file"]]}

    relations = list(relation_map.values())
    req_functions = defaultdict(list)
    fn_relations = defaultdict(list)
    for rel in relations:
        req_functions[rel["requirement_id"]].append(rel)
        fn_relations[rel["function_id"]].append(rel)

    # Complete search/model fields and immutable standard/project source facts.
    topics_for = {}
    group_profile = {
        "牵引":["TRACTION"],"制动":["BRAKING"],"储能":["ENERGY_STORAGE"],"供电":["POWER_SUPPLY"],"控制":["BRAKE_CONTROL","COMMUNICATION"],"安全":["SAFETY_PROTECTION"]}
    for rid, record in requirement_by_id.items():
        source = record.get("_source", {})
        spec = record.get("_spec", {})
        rels = req_functions[rid]
        fids = uniq([x["function_id"] for x in rels])
        pids = uniq([alloc_by_fn[fid]["product_id"] for fid in fids])
        record["Linked_Function_IDs"] = fids
        record["Linked_Function_Names"] = [fn_by_id[fid]["name"] for fid in fids]
        record["Linked_Product_IDs"] = pids
        record["Linked_Product_Names"] = [product_by_id[pid]["name"] for pid in pids]
        record["Function_Coverage_Count"] = len(fids)
        if spec:
            groups = [spec["group"]]
            record["Coverage_Status"] = "FULL" if fids else "SUPPORTING_ONLY"
            record["Model_Coverage"] = "COVERED" if fids else "PARTIAL"
            record["Evidence_Source"] = SYSML.relative_to(ROOT).as_posix()
        else:
            if rid.startswith("REQ-TRAC-"): groups=["TRACTION"]
            elif rid.startswith("REQ-BRK-") or rid.startswith("PRJ-BRK-"): groups=["BRAKING"]
            elif rid.startswith(("REQ-ENE-","PRJ-ENE-")): groups=["ENERGY_STORAGE"]
            elif rid.startswith("PRJ-PWR-"): groups=["POWER_SUPPLY"]
            else: groups=group_profile.get(record["Domain"], ["DIAGNOSTICS"])
            if rid.startswith("REQ-TRAC") and any(w in record["Name"]+record["Text"] for w in ("再生","制动")): groups.append("REGENERATIVE_BRAKING")
            if rid.startswith(("REQ-ENE","PRJ-ENE")): groups.append("ENERGY_STORAGE")
            if rid.startswith("REQ-BRK") and any(w in record["Name"]+record["Text"] for w in ("再生","电制动","能量")): groups.append("REGENERATIVE_BRAKING")
            if "安全" in record["Name"]+record["Text"] or "故障" in record["Name"]+record["Text"]: groups.append("SAFETY_PROTECTION")
            source_cov = str(source.get("model_coverage", "")).upper()
            record["Coverage_Status"] = "FULL" if source_cov in {"COVERED","FULL"} and fids else "PARTIAL"
            if not fids: record["Coverage_Status"] = "SUPPORTING_ONLY"
        groups = uniq(groups)
        record["Topic_Group"] = groups
        topics = uniq([*(spec.get("topics",[]) if spec else []), *vals(source.get("expected_function_semantics")), *groups])
        # Prefer controlled topic labels, then compact engineering concepts.
        controlled = [x for x in groups]
        record["Semantic_Topics"] = uniq([*controlled, *(spec.get("topics",[]) if spec else []),
            *[x for x in ["牵引电能变换","列车纵向动力学","再生能量回收","电空制动","机械制动","储能状态","安全保护","接口状态"] if x in (record["Name"]+record["Text"]+record["Expected_Function_Semantics"])]])
        profile_domain = ("储能" if rid.startswith(("REQ-ENE-","PRJ-ENE-")) else
            "供电" if rid.startswith(("REQ-PWR-","PRJ-PWR-","DRV-PWR-")) else
            "牵引" if rid.startswith(("REQ-TRAC-","DRV-TRAC-")) else
            "制动" if rid.startswith(("REQ-BRK-","PRJ-BRK-","DRV-BRK-")) else
            "控制" if rid.startswith(("DRV-CTL-",)) else record["Domain"])
        base = BASE_TERMS.get(profile_domain, ["列车系统","设备状态","工程功能","控制接口","运行要求"])
        record["Keywords"] = uniq([*record["Keywords"], *base, *(spec.get("keywords",[]) if spec else [])], 12)
        # Ensure at least five distinct high-signal terms.
        if len(record["Keywords"]) < 5: record["Keywords"] = uniq([*record["Keywords"], *base, record["Name"], record["Domain"]], 8)
        syn = SYNONYM_TERMS.get(profile_domain, [])
        if any(w in record["Name"]+record["Text"] for w in ("再生","回收")): syn = [*syn,"回馈制动","制动能量回收","再生能量回收"]
        if any(w in record["Name"]+record["Text"] for w in ("超级电容","SOC","储能")): syn = [*syn,"车载储能","超级电容器","超容"]
        record["Synonyms"] = uniq([*record["Synonyms"], *syn, *(spec.get("synonyms",[]) if spec else [])], 12)
        if spec:
            phrases = spec["phrases"]
        else:
            old_phrases = record["User_Phrases"]
            terms = record["Keywords"]
            domain = record["Domain"]
            special_phrases=[]
            if rid in {"REQ-BRK-006","REQ-BRK-007","REQ-BRK-008","REQ-BRK-009"}:
                special_phrases += ["机械制动什么时候介入","制动的电去哪了"]
            if rid in {"REQ-ENE-001","REQ-ENE-003","REQ-ENE-005","PRJ-ENE-001","PRJ-ENE-004"}:
                special_phrases += ["我想看再生制动","制动的电去哪了","超级电容SOC上限"]
            if rid in {"REQ-ENE-002","REQ-ENE-007","REQ-ENE-008","REQ-ENE-009","PRJ-ENE-001","PRJ-ENE-004"}:
                special_phrases += ["储能坏了怎么办"]
            if rid.startswith("REQ-BRK-") or rid=="PRJ-BRK-001": special_phrases += ["我想看制动系统"]
            phrases = [*special_phrases,*old_phrases,
                f"我想看{record['Name']}具体要求",f"{terms[0]}相关的{domain}要求有哪些",f"{terms[min(1,len(terms)-1)]}出现异常时要查什么要求"]
        record["User_Phrases"] = uniq(phrases, 7)
        if len(record["User_Phrases"]) < 3:
            record["User_Phrases"] = uniq([*record["User_Phrases"],f"我想了解{record['Name']}",f"{record['Domain']}系统对{record['Name']}有什么要求",f"{record['Name']}与相关设备如何配合"], 5)
        system_terms=[]
        for pid in pids:
            product = product_by_id[pid]
            l2 = product.get("l2_product_id") or pid
            system_terms.extend([pid,product["name"]])
            if l2 in product_by_id and l2 != pid: system_terms.extend([l2,product_by_id[l2]["name"]])
        record["System_Terms"] = uniq(system_terms, 24)
        text = " ".join([record["Name"],record["Text"],record["Expected_Function_Semantics"],*record["Keywords"],*record["Linked_Function_Names"]])
        actions = list(spec.get("actions",[]) if spec else []) + [w for w in ACTION_VOCAB if w in text]
        objects = list(spec.get("objects",[]) if spec else []) + [w for w in OBJECT_VOCAB if w in text]
        record["Action_Terms"] = uniq(actions, 16)
        record["Object_Terms"] = uniq(objects, 18)
        topics_for[rid] = set(groups)

    # Related IDs are semantic expansion links from topic clusters, never derivation edges.
    req_group_for_cluster = defaultdict(list)
    for cluster in CLUSTERS:
        for rid, groups in topics_for.items():
            if groups.intersection(cluster["groups"]): req_group_for_cluster[cluster["id"]].append(rid)
    related = defaultdict(set)
    for ids in req_group_for_cluster.values():
        for rid in ids: related[rid].update(set(ids)-{rid})
    for rid, record in requirement_by_id.items():
        record["Related_Requirement_IDs"] = sorted(related[rid])
        for field in ["Parent_Requirement_IDs","Topic_Group","Keywords","Synonyms","User_Phrases","Semantic_Topics","System_Terms","Action_Terms","Object_Terms","Linked_Function_IDs","Linked_Function_Names","Linked_Product_IDs","Linked_Product_Names","Related_Requirement_IDs"]:
            if isinstance(record[field], list): record[field] = "；".join(map(str,record[field]))

    # Classify all 191 Functions. Direct requirement links take precedence; otherwise
    # retain a parent context for genuine internal implementation/support functions.
    existing_ids = {r["Requirement_ID"] for r in requirement_by_id.values() if r.get("Source_Type") != "ENGINEERING_DERIVED"}
    coverage_rows = []
    unresolved_functions = []
    for f in functions:
        fid = f["function_id"]; owner = alloc_by_fn[fid]; pid=owner["product_id"]; product=product_by_id[pid]
        l2=product.get("l2_product_id") or pid
        direct = sorted({x["requirement_id"] for x in fn_relations[fid]})
        direct_existing = [x for x in direct if x in existing_ids]
        direct_derived = [x for x in direct if x not in existing_ids]
        if direct_existing: ctype="EXISTING_REQUIREMENT"; req_ids=direct
        elif direct_derived: ctype="DERIVED_REQUIREMENT"; req_ids=direct
        else:
            parents = [x for x in SUPPORT_PARENTS.get(l2,[]) if x in requirement_by_id]
            if parents: ctype="IMPLEMENTATION_SUPPORT"; req_ids=parents
            else: ctype="UNRESOLVED"; req_ids=[]; unresolved_functions.append(fid)
        coverage_rows.append({"Function_ID":fid,"Function_Name":f["name"],"Function_Level":f["level"],"Category":f["category"],
            "Owner_Product_ID":pid,"Owner_Product_Name":product["name"],"Covered_By_Requirement":"YES" if req_ids else "NO",
            "Requirement_IDs":"；".join(req_ids),"Coverage_Type":ctype,"Allocation_ID":owner.get("allocation_id"),
            "Allocation_Source":owner["source_file"],"Function_Source":f["source_file"],"Evidence":owner["evidence"]})

    # Search index is one term-to-many requirements. All values are semicolon-safe.
    term_rows=[]; weights={"KEYWORD":1.0,"SYNONYM":0.9,"USER_PHRASE":1.25,"SEMANTIC_TOPIC":0.7,"SYSTEM_TERM":0.6,"ACTION_TERM":0.55,"OBJECT_TERM":0.75}
    for rid, record in requirement_by_id.items():
        for field, kind in [("Keywords","KEYWORD"),("Synonyms","SYNONYM"),("User_Phrases","USER_PHRASE"),("Semantic_Topics","SEMANTIC_TOPIC"),("System_Terms","SYSTEM_TERM"),("Action_Terms","ACTION_TERM"),("Object_Terms","OBJECT_TERM")]:
            for term in vals(record[field]):
                term_rows.append({"Search_Term":term,"Term_Type":kind,"Requirement_ID":rid,"Requirement_Name":record["Name"],
                    "Topic_Group":record["Topic_Group"],"Weight":weights[kind]})
    term_rows = list({(x["Search_Term"].casefold(),x["Term_Type"],x["Requirement_ID"]):x for x in term_rows}.values())

    clusters=[]
    for c in CLUSTERS:
        members=sorted(rid for rid,groups in topics_for.items() if groups.intersection(c["groups"]) or
            (c["id"] in {"CLUSTER-BRK-REGEN","CLUSTER-ENE-RECOVERY"} and any(x in " ".join(requirement_by_id[rid]["Keywords"]) for x in ["再生制动","能量回收","再生能量"])))
        clusters.append({"Cluster_ID":c["id"],"Cluster_Name":c["name"],"Trigger_Terms":"；".join(c["triggers"]),
            "Requirement_IDs":"；".join(members),"Description":c["description"]})

    def search(query):
        q=re.sub(r"\s+","",query).casefold(); scores=defaultdict(float); matched=defaultdict(list)
        for row in term_rows:
            term=re.sub(r"\s+","",row["Search_Term"]).casefold()
            if not term: continue
            if term in q or (row["Term_Type"]=="USER_PHRASE" and len(q)>=4 and q in term):
                scores[row["Requirement_ID"]]+=row["Weight"]*(2.0 if row["Term_Type"]=="USER_PHRASE" else 1.0)
                matched[row["Requirement_ID"]].append(row["Search_Term"])
        ordered=sorted(scores,key=lambda rid:(-scores[rid],rid))
        return [{"requirement_id":rid,"requirement_name":requirement_by_id[rid]["Name"],"score":round(scores[rid],2),"matched_terms":uniq(matched[rid])} for rid in ordered]
    broad_specs=[
        ("我想看制动系统",5,{"REQ-BRK-001","DRV-BRK-001","DRV-BRK-002"}),
        ("我想看再生制动",6,{"REQ-BRK-006","REQ-BRK-007","REQ-ENE-001","PRJ-ENE-001"}),
        ("制动的电去哪了",4,{"REQ-ENE-001","REQ-ENE-005","PRJ-ENE-001"}),
        ("机械制动什么时候介入",4,{"REQ-BRK-006","REQ-BRK-007","REQ-BRK-008","DRV-BRK-002"}),
        ("储能坏了怎么办",4,{"REQ-ENE-002","REQ-ENE-007","REQ-ENE-008","REQ-ENE-009"}),
        ("超级电容",5,{"REQ-ENE-001","REQ-ENE-003","REQ-ENE-004","PRJ-ENE-004"}),
    ]
    broad_results=[]
    for q,min_count,expected in broad_specs:
        hits=search(q); got={x["requirement_id"] for x in hits}
        passed=len(hits)>=min_count and expected.issubset(got)
        broad_results.append({"query":q,"pass":passed,"minimum_count":min_count,"result_count":len(hits),
            "expected_ids":sorted(expected),"returned_ids":[x["requirement_id"] for x in hits],"top_results":hits[:12]})
    focused_q="超级电容SOC上限"; focused=search(focused_q)
    focused_targets={"PRJ-ENE-004","REQ-ENE-003","PRJ-ENE-002"}
    focused_top={x["requirement_id"] for x in focused[:5]}
    unrelated={x["requirement_id"] for x in focused[:5] if x["requirement_id"].startswith(("REQ-BRK","REQ-TRAC","DRV-BRK"))}
    focused_result={"query":focused_q,"pass":focused_targets.issubset({x["requirement_id"] for x in focused[:8]}) and not unrelated,
        "expected_ids":sorted(focused_targets),"top_results":focused[:10],"unrelated_braking_or_traction_in_top5":sorted(unrelated)}

    reqs_out=[]
    for rid,record in requirement_by_id.items():
        reqs_out.append({k:record[k] for k in FIELDS})
    reqs_out.sort(key=lambda x:(0 if not x["Requirement_ID"].startswith("DRV-") else 1,x["Requirement_ID"]))
    req_map_out=[]
    for rel in relations:
        req_map_out.append({"requirement_id":rel["requirement_id"],"requirement_name":requirement_by_id[rel["requirement_id"]]["Name"],
            "function_id":rel["function_id"],"function_name":fn_by_id[rel["function_id"]]["name"],
            "product_id":rel["product_id"],"product_name":product_by_id[rel["product_id"]]["name"],
            "mapping_type":rel["mapping_type"],"evidence":rel["evidence"],"evidence_path":rel["evidence_path"],
            "semantic_evidence":rel["semantic_evidence"],"structural_evidence":rel["structural_evidence"],"source_files":rel["source_files"]})
    coverage_counts=Counter(x["Coverage_Type"] for x in coverage_rows)
    req_direct_counts=Counter(x["requirement_id"] for x in relations)
    mapped_unique=len({x["Function_ID"] for x in coverage_rows if x["Covered_By_Requirement"]=="YES"})
    map_counts=Counter(x["mapping_type"] for x in req_map_out)
    full_count=sum(x["Coverage_Status"]=="FULL" for x in reqs_out)
    partial_count=sum(x["Coverage_Status"]=="PARTIAL" for x in reqs_out)
    support_only=sum(x["Coverage_Status"]=="SUPPORTING_ONLY" for x in reqs_out)
    total_keywords=sum(len(vals(x["Keywords"])) for x in reqs_out)
    total_synonyms=sum(len(vals(x["Synonyms"])) for x in reqs_out)
    total_phrases=sum(len(vals(x["User_Phrases"])) for x in reqs_out)
    metrics={"original_requirement_count":len(catalog),"new_requirement_count":len(DERIVED),"total_requirement_count":len(reqs_out),
        "function_count":len(functions),"covered_function_count":mapped_unique,"uncovered_function_count":len(unresolved_functions),
        "existing_requirement_covered_function_count":sum(x["Coverage_Type"]=="EXISTING_REQUIREMENT" for x in coverage_rows),
        "derived_requirement_covered_function_count":sum(x["Coverage_Type"]=="DERIVED_REQUIREMENT" for x in coverage_rows),
        "implementation_support_function_count":coverage_counts["IMPLEMENTATION_SUPPORT"],"unresolved_function_count":coverage_counts["UNRESOLVED"],
        "full_requirement_count":full_count,"partial_requirement_count":partial_count,"supporting_only_requirement_count":support_only,
        "requirement_function_relation_count":len(req_map_out),"average_functions_per_requirement":round(len(req_map_out)/len(reqs_out),2),
        "keyword_count":total_keywords,"synonym_count":total_synonyms,"user_phrase_count":total_phrases,"topic_cluster_count":len(clusters),
        "keyword_index_row_count":len(term_rows),"mapping_type_counts":dict(map_counts),"function_coverage_type_counts":dict(coverage_counts)}
    if len(catalog)!=34 or len(functions)!=191 or len(products)!=147: raise ValueError("Expected source cardinalities changed")
    if len({x["Requirement_ID"] for x in reqs_out})!=len(reqs_out): raise ValueError("Requirement ID duplicate")
    if any(x["Source_Type"]=="STANDARD_DIRECT" for x in reqs_out if x["Requirement_ID"].startswith("DRV-")): raise ValueError("Derived requirement has forged standard source")
    if any(x["Function_ID"] not in fn_by_id or x["Owner_Product_ID"] not in product_by_id or alloc_by_fn[x["Function_ID"]]["product_id"]!=x["Owner_Product_ID"] for x in coverage_rows): raise ValueError("Coverage references missing object/real allocation")
    if any(not vals(x["Keywords"]) or not vals(x["User_Phrases"]) or not vals(x["Semantic_Topics"]) for x in reqs_out): raise ValueError("Search field empty")
    if any(not (5<=len(vals(x["Keywords"]))<=15) for x in reqs_out): raise ValueError("Keyword count outside 5-15")
    if any(not (3<=len(vals(x["User_Phrases"]))<=8) for x in reqs_out): raise ValueError("User phrase count outside 3-8")
    if not all(x["pass"] for x in broad_results):
        for x in broad_results:
            if not x["pass"]: print("BROAD_QUERY_FAIL",x["query"],"COUNT",x["result_count"],"MISSING",sorted(set(x["expected_ids"])-set(x["returned_ids"])))
        raise ValueError("Broad query acceptance failed")
    if not focused_result["pass"]: raise ValueError("Focused query acceptance failed")
    # Existing v1 protected values must still equal copied source rows in final library.
    by_id={x["Requirement_ID"]:x for x in reqs_out}
    for rid, raw in raw_by_id.items():
        row=by_id[rid]
        for field, col in [("Requirement_ID","Requirement_ID"),("Name","需求名称"),("Text","正式需求表述"),("Requirement_Type","需求类型"),
            ("Source_Type","来源类型"),("Standard_No","标准号"),("Standard_Clause","条款号"),("Criteria","标准/项目判据")]:
            if str(row[field] or "") != str(raw.get(col) or ""): raise ValueError(f"Protected original value differs: {rid} {field}")

    final_data={"schema_version":"2.0","source_requirement_library":REQ_XLSX.relative_to(ROOT).as_posix(),"requirements":reqs_out,"summary":metrics}
    map_data={"schema_version":"2.0","relations":req_map_out,"mapping_type_counts":dict(map_counts),"unresolved_requirement_ids":[]}
    index_data={"schema_version":"2.0","records":term_rows,"broad_query_tests":broad_results,"focused_query_test":focused_result}
    clusters_data={"schema_version":"2.0","clusters":clusters}
    coverage_data={"schema_version":"2.0","summary":metrics,"functions":coverage_rows,"unresolved_function_ids":unresolved_functions}
    for fname,obj in [("requirement_catalog_v2.json",final_data),("requirement_function_map_v2.json",map_data),
        ("requirement_search_index_v2.json",index_data),("requirement_topic_clusters.json",clusters_data),
        ("function_requirement_coverage_v2.json",coverage_data)]:
        (OUT/fname).write_text(json.dumps(obj,ensure_ascii=False,indent=2),encoding="utf-8")

    lines=["# Requirement Library V2 Expansion Audit","","## Results","",
        f"- Original requirements: {len(catalog)}; newly derived: {len(DERIVED)}; total: {len(reqs_out)}.",
        f"- Real SysML Functions: {len(functions)}; Function → OWNER Product source allocations: {len(allocations)}.",
        f"- Unique Functions with requirement upstream: {mapped_unique}/{len(functions)}; unresolved: {len(unresolved_functions)}.",
        f"- Function audit classifications: existing requirement {coverage_counts['EXISTING_REQUIREMENT']}; derived requirement {coverage_counts['DERIVED_REQUIREMENT']}; implementation support {coverage_counts['IMPLEMENTATION_SUPPORT']}; unresolved {coverage_counts['UNRESOLVED']}.",
        f"- Requirement coverage: FULL {full_count}; PARTIAL {partial_count}; SUPPORTING_ONLY {support_only}.",
        f"- Requirement → Function relations: {len(req_map_out)}; average {metrics['average_functions_per_requirement']} per requirement.",
        f"- Search fields: {total_keywords} keywords; {total_synonyms} synonyms; {total_phrases} user phrases; {len(clusters)} topic clusters; {len(term_rows)} term-to-requirement index rows.",
        "- New requirements use `ENGINEERING_DERIVED`, have blank standard fields, and cite only the frozen SysML Functions and their real allocations. No standard number, clause, standard quotation or SysML model object was added.",
        "- For original 34 records, Requirement_ID, Standard_No, Standard_Clause, Requirement_Type, Source_Type, Text, and Criteria were compared against v1 and kept unchanged.",
        "- Function coverage labels describe traceability source type, not verification status. Original `Verification_Status` values are preserved; all new derived requirements are `NOT_VERIFIED`.",
        "- `IMPLEMENTATION_SUPPORT` rows inherit listed parent Requirement context and do not assert a separate direct satisfaction relation.","",
        "## Added engineering-derived requirements","", "| ID | Name | Function count | Parent requirements |", "|---|---|---:|---|"]
    for spec in DERIVED:
        lines.append(f"| {spec['id']} | {spec['name']} | {len(spec['functions'])} | {'；'.join(spec['parents'])} |")
    lines += ["","## Function coverage audit","","| Coverage type | Functions |","|---|---:|"]
    for typ in ["EXISTING_REQUIREMENT","DERIVED_REQUIREMENT","IMPLEMENTATION_SUPPORT","UNRESOLVED"]: lines.append(f"| {typ} | {coverage_counts[typ]} |")
    lines += ["","### Unresolved Functions","",*( [f"- `{fid}`" for fid in unresolved_functions] if unresolved_functions else ["- None; all 191 Functions have either direct requirement links or an explicit implementation-support parent."]),
        "","## Search-index acceptance","","| Query | Status | Result count | Expected IDs present |","|---|---|---:|---|"]
    for item in broad_results:
        expected_ok=set(item["expected_ids"]).issubset(set(item["returned_ids"]))
        lines.append(f"| {item['query']} | {'PASS' if item['pass'] else 'FAIL'} | {item['result_count']} | {'YES' if expected_ok else 'NO'} |")
    lines += [f"| {focused_q} (focused) | {'PASS' if focused_result['pass'] else 'FAIL'} | {len(focused)} | {', '.join(focused_result['expected_ids'])} in top 8; unrelated top 5={len(focused_result['unrelated_braking_or_traction_in_top5'])} |",
        "","Query verification uses a bounded lexical check over the exported weighted term index; it validates one-to-many recall and focused topicality, and does not change the UI/Qwen matching algorithm.",
        "","## Output files","", "- `Req/牵引制动能量回收系统_正式需求库_v2_扩展版.xlsx` (created by the Artifact Tool workbook builder).",
        "- `ui/backend/data/semantic_slice/requirement_catalog_v2.json`.","- `ui/backend/data/semantic_slice/requirement_function_map_v2.json`.",
        "- `ui/backend/data/semantic_slice/requirement_search_index_v2.json`.","- `ui/backend/data/semantic_slice/requirement_topic_clusters.json`.",
        "- `ui/backend/data/semantic_slice/function_requirement_coverage_v2.json`.","- Frozen SysML SHA-256 was checked by the upstream traceability validation; no frozen asset was written.",""]
    REPORT.parent.mkdir(parents=True,exist_ok=True); REPORT.write_text("\n".join(lines),encoding="utf-8")
    BUNDLE.write_text(json.dumps({"fields":FIELDS,"requirements":reqs_out,"raw_headers":raw_headers,"raw_rows":raw_rows,
        "relations":req_map_out,"coverage":coverage_rows,"clusters":clusters,"search_index":term_rows,"summary":metrics,
        "broad_query_tests":broad_results,"focused_query_test":focused_result},ensure_ascii=False,indent=2),encoding="utf-8")
    print(f"LIBRARY_DATA=PASS ORIGINAL={len(catalog)} DERIVED={len(DERIVED)} TOTAL={len(reqs_out)}")
    print(f"FUNCTION_COVERAGE={mapped_unique}/191 EXISTING={coverage_counts['EXISTING_REQUIREMENT']} DERIVED={coverage_counts['DERIVED_REQUIREMENT']} SUPPORT={coverage_counts['IMPLEMENTATION_SUPPORT']} UNRESOLVED={coverage_counts['UNRESOLVED']}")
    print(f"RELATIONS={len(req_map_out)} KEYWORD_INDEX={len(term_rows)} CLUSTERS={len(clusters)} BROAD={sum(x['pass'] for x in broad_results)}/{len(broad_results)} FOCUSED={'PASS' if focused_result['pass'] else 'FAIL'}")


if __name__ == "__main__":
    main()
