from __future__ import annotations

import hashlib
import json
import re
import sys
from collections import Counter, defaultdict
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook

BACKEND = Path(__file__).resolve().parents[1]
ROOT = BACKEND.parents[1]
sys.path.insert(0, str(BACKEND))

from app.services.model_graph_service import ModelGraphService, attrs, block_at  # noqa: E402

SOURCE_XLSX = ROOT / "Req/牵引制动能量回收系统_正式需求库_v1.xlsx"
SOURCE_SYSML = ROOT / "work/sysmlv2/full_engineering_model/01_generated/Rail_MBSE_Full_v1.sysml"
OUT = BACKEND / "data/semantic_slice"
REPORTS = ROOT / "work/reports"
TRACE_KEYS = ("sysml_refs", "interface_refs", "fmi_refs", "fmu_refs", "ssp_refs", "simulation_refs", "vv_refs")

FIELD_MAP = {
    "Requirement_ID": "requirement_id", "领域": "domain", "需求名称": "name", "正式需求表述": "text",
    "需求类型": "requirement_type", "来源类型": "source_type", "标准号": "standard_no", "条款号": "standard_clause",
    "标准/项目判据": "criteria", "适用条件": "applicable_conditions", "关键词": "keywords", "同义词": "synonyms",
    "典型用户问法": "user_phrases", "预期功能语义": "expected_function_semantics", "已知Function_ID": "known_function_id",
    "主要Product": "primary_product", "模型覆盖": "model_coverage", "当前V&V状态": "verification_status",
    "备注": "notes", "证据URL": "evidence_url",
}

ALIAS_L2 = {
    "牵引变流器": "5100", "牵引电机": "5300", "主变压器": "4500", "接触网": "4100", "线路": "4100",
    "制动协调": "7200", "机械制动": "7200", "空气制动": "7200", "制动系统": "7200",
    "监测/控制系统": "8100", "列车控制": "8100", "车载再生储能": "X100", "X100储能": "X100", "DC/DC": "X100",
}
STOP_TERMS = {"需求", "系统", "功能", "状态", "控制", "计算", "生成", "输出", "输入", "实现", "提供", "执行", "检测", "传递"}


def dump(name: str, value):
    (OUT / name).write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding="utf-8")


def split_values(value):
    if value is None: return []
    return [x.strip() for x in re.split(r"[;；,，\n]+", str(value)) if x and x.strip()]


def clean(value):
    return "" if value is None else str(value).strip()


def traceability(**values):
    return {k: list(values.get(k, [])) for k in TRACE_KEYS}


def terms(value):
    values = split_values(value)
    result = set()
    for item in values:
        compact = re.sub(r"[^0-9A-Za-z\u4e00-\u9fff]+", "", item)
        if len(compact) >= 2 and compact not in STOP_TERMS: result.add(compact)
        result.update(x for x in re.findall(r"[\u4e00-\u9fff]{2,6}|[A-Za-z]+\d*", item) if x not in STOP_TERMS)
    return result


def bigrams(value):
    text = re.sub(r"[^0-9a-z\u4e00-\u9fff]", "", value.lower())
    return {text[i:i + 2] for i in range(max(0, len(text) - 1))}


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def read_requirements():
    workbook = load_workbook(SOURCE_XLSX, read_only=True, data_only=True)
    sheet_counts = {}
    all_rows = {}
    for sheet in workbook.worksheets:
        rows = [list(row) for row in sheet.iter_rows(values_only=True) if any(v is not None for v in row)]
        sheet_counts[sheet.title] = max(0, len(rows) - 1)
        all_rows[sheet.title] = rows
    rows = all_rows["正式需求库"]
    headers = [clean(x) for x in rows[0]]
    requirements = []
    for values in rows[1:]:
        source = {headers[i]: values[i] if i < len(values) else None for i in range(len(headers))}
        req = {target: clean(source.get(original)) for original, target in FIELD_MAP.items()}
        for key in ("keywords", "synonyms", "user_phrases"):
            req[key] = split_values(req[key])
        req["traceability"] = traceability()
        requirements.append(req)
    codex_rows = all_rows["Codex导入视图"]
    codex_ids = {clean(row[0]) for row in codex_rows[1:] if row and row[0]}
    return requirements, sheet_counts, codex_ids


def read_model():
    graph = ModelGraphService(ROOT).get_full_model_graph()
    products = []
    product_by_code = {}
    for item in graph["nodes"]:
        p = item["data"]
        product = {"product_id": p["code"], "name": p["display_name"], "description": "", "level": p["level"],
            "parent_product_id": p.get("parent_code"), "l2_product_id": p.get("l2_code"), "source": "SysML v2 part def",
            "source_file": p["source_path"], "function_ids": [], "traceability": traceability(sysml_refs=[p["sysml_id"]])}
        products.append(product); product_by_code[p["code"]] = product
    text = SOURCE_SYSML.read_text(encoding="utf-8")
    line_starts = [0] + [m.end() for m in re.finditer("\n", text)]
    import bisect
    functions = []
    allocations = defaultdict(list)
    for m in re.finditer(r"allocation\s+(\w+)\s+allocate\s+fu_(F_\w+)\s+to\s+ap_(?:N_)?([A-Z0-9]+)\s*;", text):
        allocations["FN_" + m[2]].append((m[1], m[3], bisect.bisect_right(line_starts, m.start())))
    function_product = []
    for m in re.finditer(r"action def\s+(FN_F_\w+)\s*\{", text):
        a = attrs(block_at(text, m.start())); fid = m[1]
        links = [(aid, code, line) for aid, code, line in allocations.get(fid, []) if code in product_by_code]
        if not links: continue
        src_line = bisect.bisect_right(line_starts, m.start())
        function = {"function_id": fid, "name": a.get("chineseName") or fid, "description": a.get("englishName") or "",
            "category": a.get("category"), "level": a.get("functionLevel"), "source": "SysML v2 action def",
            "source_file": f"{SOURCE_SYSML.relative_to(ROOT).as_posix()}:{src_line}", "product_ids": [x[1] for x in links],
            "keywords": [], "aliases": [], "traceability": traceability(sysml_refs=[fid])}
        function["source_classification"], function["classification_basis"] = function_kind(function)
        functions.append(function)
        for allocation_id, code, line in links:
            relation = {"function_id": fid, "product_id": code, "allocation_id": allocation_id, "relation_type": "ALLOCATED_TO", "mapping_source": "SYSML",
                "source_file": f"{SOURCE_SYSML.relative_to(ROOT).as_posix()}:{line}",
                "evidence": f"allocation {allocation_id} allocate {fid.removeprefix('FN_')} to {code}", "status": "RESOLVED"}
            function_product.append(relation); product_by_code[code]["function_ids"].append(fid)
    return graph, functions, products, function_product


def resolve_product_codes(value, products):
    hits = set(re.findall(r"(?<![A-Z0-9])(?:[3-8]\d{3}|[DEX]\d{3})(?![A-Z0-9])", value or ""))
    for token in split_values(value):
        for alias, code in ALIAS_L2.items():
            if alias in token: hits.add(code)
        exact = [p["product_id"] for p in products if p["name"] and (p["name"] in token or token in p["name"])]
        if len(exact) == 1: hits.add(exact[0])
    return hits


def make_requirement_function(requirements, functions, products):
    fn_by_id = {f["function_id"]: f for f in functions}
    relations = []; suggestions = []
    for req in requirements:
        explicit = split_values(req["known_function_id"])
        for fid in explicit:
            if fid in fn_by_id:
                relations.append({"requirement_id": req["requirement_id"], "function_id": fid, "relation_type": "SATISFIED_BY",
                    "confidence": 1.0, "mapping_source": "EXPLICIT", "source_file": str(SOURCE_XLSX.relative_to(ROOT)).replace("\\", "/"),
                    "evidence": f"正式需求库 known_function_id={fid}", "notes": "工程映射；confidence 不是V&V结论"})
        # Lexical similarity is not traceability evidence. Model-derived links are
        # added by build_requirement_strategy.py after resolving real model paths.
        if explicit: continue
    return relations, suggestions


def function_kind(function):
    name = function["name"]
    if re.search(r"(向.{1,12}(传递|输出|发送|输送)|给.{1,12}(提供|输出)|从.{1,12}(接收|获取))", name):
        return "FLOW_DERIVED_DESCRIPTION", "名称符合方向性传递模板；仅为名称分类，需对照真实connector核查"
    if function.get("category") == "PHYSICAL_TRANSPORT":
        return "INTERACTION_FUNCTION", "原始SysML action def category=PHYSICAL_TRANSPORT"
    try:
        if int(function.get("level") or 0) <= 2:
            return "SYSTEM_FUNCTION", "原始SysML action def functionLevel<=2"
        return "ACTION", "原始SysML action def functionLevel>=3"
    except (TypeError, ValueError):
        return "UNKNOWN", "SysML action def未声明可解析的functionLevel"


def build_interaction_map(functions, function_product, graph):
    owners = defaultdict(list)
    for relation in function_product:
        owners[relation["function_id"]].append({"product_id": relation["product_id"],
            "role": "OWNER", "source": "SYSML_ALLOCATION", "source_file": relation["source_file"],
            "evidence": relation["evidence"]})
    interactions = []
    model_text = SOURCE_SYSML.read_text(encoding="utf-8")
    connector_interface = {m.group(1): m.group(2) for m in re.finditer(
        r"interface\s+c_(CG_[A-F0-9]+)\s*:\s*(IF_IF_\w+)\s+connect", model_text)}
    for interface in graph["aggregated_interfaces"]:
        if interface["level"] != 2:
            continue
        source_code = interface["source_node_id"].split("::", 1)[1]
        target_code = interface["target_node_id"].split("::", 1)[1]
        connectors = [x.split("::", 1)[1] for x in interface.get("leaf_connections", []) if "::" in x]
        leaf_refs = list(interface.get("leaf_connections", []))
        for flow in interface.get("flow_items", []):
            label = f"{flow.get('display_name', '')} {flow.get('raw_name', '')} {flow.get('kind', '')}".lower()
            kind = "unknown"
            if flow.get("kind") == "state": kind = "state"
            elif any(x in label for x in ("cmd", "control", "command", "指令", "控制", "信号", "状态")): kind = "signal"
            elif any(x in label for x in ("force", "brakeforce", "制动力", "牵引力")): kind = "force"
            elif any(x in label for x in ("power", "energy", "功率", "能量", "电压", "电流", "转矩", "转速")): kind = "energy"
            elif flow.get("kind") == "physical": kind = "unknown"
            elif flow.get("kind") == "data": kind = "signal"
            interactions.append({"interaction_id": interface["id"] + "::" + flow["id"],
                "source_product_id": source_code, "target_product_id": target_code,
                "flow_name": flow.get("display_name") or flow.get("raw_name"), "flow_kind": kind,
                "flow_ref": flow["id"], "interface_refs": sorted({connector_interface[c] for c in connectors if c in connector_interface}),
                "aggregate_interface_ref": interface["id"], "leaf_connection_refs": leaf_refs,
                "connector_refs": connectors, "source": "TASK_SLICE_STAGE4",
                "evidence": f"真实L2聚合接口 {interface['id']}；SysML叶级连接 {', '.join(connectors)}",
                "source_path": interface.get("source_path"), "raw_flow_kind": flow.get("kind")})
    per_function = []
    for function in functions:
        # Product adjacency alone cannot prove a specific action participates in a connector.
        per_function.append({"function_id": function["function_id"], "allocated_products": owners[function["function_id"]],
            "interactions": [], "interaction_status": "UNRESOLVED",
            "unresolved_reason": "现有action def/allocation未将此Function引用到具体port、interface或connector；未按owner邻接关系推断参与。"})
    return per_function, interactions


def main():
    OUT.mkdir(parents=True, exist_ok=True); REPORTS.mkdir(parents=True, exist_ok=True)
    generated_at = datetime.now(timezone.utc).isoformat()
    requirements, sheet_counts, codex_ids = read_requirements()
    graph, functions, products, function_product = read_model()
    requirement_function, candidate_suggestions = make_requirement_function(requirements, functions, products)
    function_interactions, l2_interactions = build_interaction_map(functions, function_product, graph)
    req_by_id = {r["requirement_id"]: r for r in requirements}; fn_by_id = {f["function_id"]: f for f in functions}; prod_by_id = {p["product_id"]: p for p in products}
    rf_by_req = defaultdict(list); fp_by_fn = defaultdict(list)
    for x in requirement_function: rf_by_req[x["requirement_id"]].append(x)
    for x in function_product: fp_by_fn[x["function_id"]].append(x)
    aggregate = {}
    for rid, req in req_by_id.items():
        fn_items = []
        for rel in rf_by_req[rid]:
            fn = fn_by_id[rel["function_id"]]
            fn_items.append({**fn, "mapping": rel, "products": [{**prod_by_id[x["product_id"]], "mapping": x} for x in fp_by_fn[fn["function_id"]]]})
        aggregate[rid] = {"requirement": req, "functions": fn_items}
    warnings = []
    req_counts = Counter(r["requirement_id"] for r in requirements); prod_counts = Counter(p["product_id"] for p in products)
    warnings += [f"duplicate requirement: {k}" for k, v in req_counts.items() if v > 1]
    warnings += [f"duplicate product: {k}" for k, v in prod_counts.items() if v > 1]
    warnings += [f"unknown function relation: {x}" for x in requirement_function if x["function_id"] not in fn_by_id]
    warnings += [f"unknown product relation: {x}" for x in function_product if x["product_id"] not in prod_by_id]
    valid_status = {"VERIFIED", "NOT_VERIFIED", "PARTIALLY_VERIFIED"}
    warnings += [f"invalid verification_status: {r['requirement_id']}={r['verification_status']}" for r in requirements if r["verification_status"] not in valid_status]
    unresolved_requirements = [r["requirement_id"] for r in requirements if not rf_by_req[r["requirement_id"]]]
    unresolved_functions = [f["function_id"] for f in functions if not fp_by_fn[f["function_id"]]]
    semantic_index = [{"requirement_id": r["requirement_id"], "search_text": " ".join([r["domain"], r["name"], r["text"],
        *r["keywords"], *r["synonyms"], *r["user_phrases"], r["expected_function_semantics"], r["primary_product"]]),
        "function_ids": [x["function_id"] for x in rf_by_req[r["requirement_id"]]],
        "product_ids": sorted({p["product_id"] for x in rf_by_req[r["requirement_id"]] for p in fp_by_fn[x["function_id"]]})} for r in requirements]
    manifest = {"generated_at": generated_at, "requirement_source": str(SOURCE_XLSX.relative_to(ROOT)).replace("\\", "/"),
        "requirement_source_sha256": sha(SOURCE_XLSX), "sysml_source_sha256": sha(SOURCE_SYSML), "sheet_counts": sheet_counts,
        "codex_import_id_count": len(codex_ids), "codex_import_ids_match": codex_ids == set(req_by_id),
        "requirement_count": len(requirements), "function_count": len(functions), "product_count": len(products),
        "requirement_function_relation_count": len(requirement_function), "candidate_suggestion_count": len(candidate_suggestions),
        "function_product_relation_count": len(function_product), "l2_interaction_count": len(l2_interactions),
        "unresolved_requirement_count": len(unresolved_requirements), "unresolved_function_count": len(unresolved_functions),
        "warnings": warnings, "files": ["requirement_catalog.json", "function_catalog.json", "product_catalog.json",
        "requirement_function_map.json", "requirement_function_candidates.json", "function_product_map.json", "function_interaction_map.json", "l2_interaction_catalog.json", "requirement_function_product_map.json", "semantic_query_examples.json", "semantic_index.json", "semantic_cache.json"]}
    dump("requirement_catalog.json", {"schema_version": "1.0", "source_file": manifest["requirement_source"], "generated_at": generated_at, "requirements": requirements})
    dump("function_catalog.json", {"schema_version": "1.0", "source_file": str(SOURCE_SYSML.relative_to(ROOT)).replace("\\", "/"), "generated_at": generated_at, "functions": functions})
    dump("product_catalog.json", {"schema_version": "1.0", "source_file": str(SOURCE_SYSML.relative_to(ROOT)).replace("\\", "/"), "generated_at": generated_at, "products": products})
    dump("requirement_function_map.json", {"relations": requirement_function, "unresolved_requirement_ids": unresolved_requirements})
    dump("requirement_function_candidates.json", {"status": "UNREVIEWED_CANDIDATES_ONLY", "relations": candidate_suggestions})
    dump("function_product_map.json", {"relations": function_product, "unresolved_function_ids": unresolved_functions})
    dump("function_interaction_map.json", {"functions": function_interactions, "unresolved_function_ids": [x["function_id"] for x in function_interactions if not x["interactions"]]})
    dump("l2_interaction_catalog.json", {"source": graph["authority"], "source_sha256": graph["source_sha256"], "interactions": l2_interactions})
    dump("requirement_function_product_map.json", aggregate)
    dump("semantic_query_examples.json", {"queries": ["我想看制动系统", "我想看再生制动", "我想看能量回收", "我只想看X100", "我想看牵引系统", "我想看制动时为什么需要机械制动", "制动产生的能量是怎么进入超级电容的", "今天天气怎么样"]})
    dump("semantic_index.json", {"generated_at": generated_at, "records": semantic_index})
    if not (OUT / "semantic_cache.json").exists(): dump("semantic_cache.json", {})
    dump("data_build_manifest.json", manifest)
    integrity = ["# Semantic Data Integrity Report", "", f"Generated: {generated_at}", "", f"- Duplicate Requirement IDs: {sum(v > 1 for v in req_counts.values())}",
        f"- Duplicate Product IDs: {sum(v > 1 for v in prod_counts.values())}", f"- Dangling Requirement→Function relations: {sum(x['function_id'] not in fn_by_id for x in requirement_function)}",
        f"- Dangling Function→Product relations: {sum(x['product_id'] not in prod_by_id for x in function_product)}", f"- Unmapped Requirements: {len(unresolved_requirements)}",
        f"- Functions without Product: {len(unresolved_functions)}", f"- Missing mapping sources: {sum(not x.get('source_file') for x in requirement_function + function_product)}", "",
        "## Unmapped Requirements", "", *(f"- {x}" for x in unresolved_requirements or ["0"]), "", "## Warnings", "", *(f"- {x}" for x in warnings or ["0"])]
    (REPORTS / "SEMANTIC_DATA_INTEGRITY_REPORT.md").write_text("\n".join(integrity), encoding="utf-8")
    print(f"SEMANTIC_DATA_BUILD = {'PASS' if not warnings else 'PASS_WITH_WARNINGS'}")
    print(f"COUNTS = {len(requirements)}/{len(functions)}/{len(products)}/{len(requirement_function)}/{len(function_product)}/{len(l2_interactions)}")
    print(f"UNRESOLVED = {len(unresolved_requirements)}/{len(unresolved_functions)}")


if __name__ == "__main__": main()
